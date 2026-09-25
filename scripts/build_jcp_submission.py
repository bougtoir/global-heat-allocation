"""Build the Journal of Cleaner Production manuscript and submission package."""

from __future__ import annotations

import html
import json
import re
import shutil
import zipfile
from pathlib import Path

import matplotlib.pyplot as plt
import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as PptxRGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches as PptxInches
from pptx.util import Pt as PptxPt

ROOT = Path(__file__).resolve().parents[1]
MANUSCRIPT_SOURCE = ROOT / "manuscript" / "jcp_manuscript.md"
SUBMISSION = ROOT / "submission" / "jcp"
FIGURE_DIR = SUBMISSION / "figures"
RESULTS = ROOT / "results"
TABLES = RESULTS / "tables"
FIGURES = RESULTS / "figures"
REFERENCE_ROOT = ROOT / "data" / "raw" / "references"
REFERENCE_REGISTRY = ROOT / "provenance" / "jcp_references_verified.csv"
MANUSCRIPT_VALUES = ROOT / "provenance" / "manuscript_values.csv"
CASE_EVIDENCE = ROOT / "provenance" / "case_specific_evidence.csv"
FAILURE_MATRIX = TABLES / "free_sink_failure_matrix.csv"
STRUCTURAL = TABLES / "structural_replication.csv"
PATHWAYS = RESULTS / "environmental_constraints" / "frontier_pathway_assessment.csv"
PRACTICAL_MODES = RESULTS / "techno_economic" / "practical_pareto_modes.csv"
WATERFALL = TABLES / "global_to_real_constraint_waterfall.csv"

TITLE = (
    "There Is No Free Heat Sink: Detecting Environmental Burden Shifting "
    "in Thermal-Energy Optimization"
)

FIGURE_SOURCES = {
    1: "figure_1_no_free_heat_sink_framework",
    2: "supplementary_figure_s4_peak_marginal_burden_map",
    3: "figure_3_progressive_sink_constraints",
    4: "figure_3_allocation_mode_comparison",
    5: "figure_5_constraint_waterfall",
    6: "figure_6_practical_modes",
}

FIGURE_CAPTIONS = {
    1: (
        "Figure 1. Conservation-aware no-free-heat-sink framework. A local "
        "human-burden objective can select an apparent receiving sink, but "
        "environmental and practical evidence gates must validate the full "
        "pathway before a cleaner-production endpoint is supported."
    ),
    2: (
        "Figure 2. Annual peak population-weighted marginal thermal-stress "
        "burden per GJ on the common global grid. The map shows heterogeneity "
        "used by the screening objective; it does not identify an "
        "environmentally safe source or receiving location."
    ),
    3: (
        "Figure 3. Progressive receiving constraints and evidence gates. "
        "Burden-space values are reported only where defined; physical-energy "
        "evidence stages remain categorical rather than being converted into "
        "burden units, MWh-th, USD, or CO2e."
    ),
    4: (
        "Figure 4. Matched-event spatial and same-location temporal screening "
        "responses. The temporal case uses the same source event and retains "
        "part of the spatial burden-space value without selecting a remote "
        "receiving cell."
    ),
    5: (
        "Figure 5. Global-to-real evidence waterfall. The theoretical "
        "burden-space and measured physical-energy branches retain separate "
        "units and meet only through explicit evidence eligibility."
    ),
    6: (
        "Figure 6. Conditional direct-reuse and storage sensitivity modes. "
        "All modes remain conditional and zero practical endpoints are "
        "supported under the current public evidence."
    ),
}


def _configure_document(document: Document, *, compact: bool = False) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11 if compact else 12)
    normal.paragraph_format.line_spacing = 1.35 if compact else 2
    normal.paragraph_format.space_after = Pt(4 if compact else 0)
    for heading in ["Title", "Heading 1", "Heading 2", "Heading 3"]:
        document.styles[heading].font.name = "Arial"
        document.styles[heading].font.color.rgb = RGBColor(0, 0, 0)
    for section in document.sections:
        margin = 0.85 if compact else 1
        section.top_margin = Inches(margin)
        section.bottom_margin = Inches(margin)
        section.left_margin = Inches(margin)
        section.right_margin = Inches(margin)


def _set_table_borders(table) -> None:
    properties = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "insideH", "bottom"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:color"), "808080")
        borders.append(border)
    for edge in ("left", "insideV", "right"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "nil")
        borders.append(border)
    properties.append(borders)


def _add_cited_paragraph(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    for part in re.split(r"(\{[^}]+\})", text):
        if part.startswith("{") and part.endswith("}"):
            run = paragraph.add_run(part[1:-1])
            run.font.superscript = True
        else:
            paragraph.add_run(part)


def _add_table(document: Document, frame: pd.DataFrame, caption: str) -> None:
    caption_paragraph = document.add_paragraph()
    caption_paragraph.paragraph_format.space_before = Pt(12)
    caption_paragraph.paragraph_format.keep_with_next = True
    caption_paragraph.add_run(caption).bold = True
    table = document.add_table(rows=1, cols=len(frame.columns))
    table.style = "Table Grid"
    _set_table_borders(table)
    for column, cell in zip(frame.columns, table.rows[0].cells, strict=True):
        cell.text = str(column)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    header_properties = table.rows[0]._tr.get_or_add_trPr()
    repeat_header = OxmlElement("w:tblHeader")
    repeat_header.set(qn("w:val"), "true")
    header_properties.append(repeat_header)
    for row in frame.itertuples(index=False):
        cells = table.add_row().cells
        for value, cell in zip(row, cells, strict=True):
            if pd.isna(value):
                cell.text = ""
            elif isinstance(value, float):
                cell.text = f"{value:.4g}"
            else:
                rendered = str(value)
                if "_" in rendered and not re.fullmatch(r"[A-Z0-9_]+", rendered):
                    rendered = rendered.replace("_", " ")
                cell.text = rendered
    for row in table.rows:
        row_properties = row._tr.get_or_add_trPr()
        row_properties.append(OxmlElement("w:cantSplit"))
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.line_spacing = 1
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(8)


def _compact_heading(column: object) -> str:
    labels = {
        "energy_gj": "Energy (GJ)",
        "distance_km": "Distance (km)",
        "source_temperature_change_k": "Source temperature change (K)",
        "sink_temperature_change_k": "Sink temperature change (K)",
        "net_benefit_per_gj": "Net benefit per GJ",
        "screening_value_burden_units_per_gj": ("Screening value (burden units/GJ)"),
        "useful_heat_delivered_mwh": "Useful heat delivered (MWh-th)",
        "residual_source_rejection_mwh": ("Residual source rejection (MWh-th)"),
    }
    key = str(column)
    if key in labels:
        return labels[key]
    heading = key.replace("_", " ").title()
    replacements = {
        " Co2E": " CO2e",
        " Cop": " COP",
        " Gj": " GJ",
        " Id": " ID",
        " Km": " km",
        " Mw": " MW",
        " Mwh": " MWh",
    }
    for old, new in replacements.items():
        heading = heading.replace(old, new)
    return heading


def _compact_value(value: object) -> str:
    if pd.isna(value):
        return ""
    if isinstance(value, float):
        return f"{value:.3g}"
    return {
        "MWh_th": "MWh-th",
        "MW_th": "MW-th",
        "burden_units_per_GJ": "burden units/GJ",
    }.get(str(value), str(value))


def _add_compact_table(
    document: Document,
    frame: pd.DataFrame,
    caption: str,
) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run(caption).bold = True
    table = document.add_table(rows=1, cols=len(frame.columns))
    table.style = "Table Grid"
    _set_table_borders(table)
    for column, cell in zip(frame.columns, table.rows[0].cells, strict=True):
        cell.text = _compact_heading(column)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    header_properties = table.rows[0]._tr.get_or_add_trPr()
    repeat_header = OxmlElement("w:tblHeader")
    repeat_header.set(qn("w:val"), "true")
    header_properties.append(repeat_header)
    for row in frame.itertuples(index=False):
        cells = table.add_row().cells
        for value, cell in zip(row, cells, strict=True):
            cell.text = _compact_value(value)
    for row in table.rows:
        for cell in row.cells:
            for cell_paragraph in cell.paragraphs:
                cell_paragraph.paragraph_format.line_spacing = 1
                cell_paragraph.paragraph_format.space_after = Pt(0)
                for run in cell_paragraph.runs:
                    run.font.size = Pt(7)


def _add_inline_figure(document: Document, number: int) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(12)
    source = FIGURES / f"{FIGURE_SOURCES[number]}.png"
    paragraph.add_run().add_picture(str(source), width=Inches(6.2))
    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.add_run(FIGURE_CAPTIONS[number]).bold = True


def _normalize_zip_archive(path: Path) -> None:
    with zipfile.ZipFile(path) as archive:
        members = [(item.filename, archive.read(item)) for item in archive.infolist()]
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for filename, data in members:
            item = zipfile.ZipInfo(filename, date_time=(1980, 1, 1, 0, 0, 0))
            item.compress_type = zipfile.ZIP_DEFLATED
            item.create_system = 3
            item.external_attr = 0o600 << 16
            archive.writestr(item, data)


def _reference_keys() -> list[str]:
    registry = pd.read_csv(REFERENCE_REGISTRY).fillna("")
    included = registry.loc[registry["jcp_use"].eq("include")].copy()
    included["jcp_reference_order"] = included["jcp_reference_order"].astype(int)
    included = included.sort_values("jcp_reference_order")
    if included["jcp_reference_order"].tolist() != list(range(1, len(included) + 1)):
        raise ValueError("JCP reference registry is not sequential.")
    if len(included) > 50:
        raise ValueError("JCP reference limit exceeded.")
    return included["citation_key"].tolist()


def _latest_reference_snapshot() -> Path:
    snapshots = sorted(REFERENCE_ROOT.glob("crossref_*"))
    if not snapshots:
        raise FileNotFoundError("No verified reference snapshot is available.")
    return snapshots[-1]


def _crossref_reference(number: int, message: dict) -> str:
    def first(values: list[str] | None) -> str:
        return values[0] if values else ""

    authors = message.get("author", [])
    names = []
    for author in authors[:6]:
        family = author.get("family", "")
        initials = "".join(part[0] for part in author.get("given", "").split() if part)
        names.append(f"{family} {initials}".strip() or author.get("name", ""))
    if len(authors) > 6:
        names.append("et al")
    title = html.unescape(first(message.get("title"))).rstrip(".")
    journal = html.unescape(
        first(message.get("container-title")) or message.get("publisher", "")
    )
    year = message.get("issued", {}).get("date-parts", [[None]])[0][0]
    volume = message.get("volume", "")
    issue = message.get("issue", "")
    pages = message.get("page", "")
    locator = volume + (f"({issue})" if issue else "")
    if pages:
        locator += f":{pages}"
    doi = message.get("DOI", "")
    if doi == "10.2172/2329591":
        return (
            f"{number}. Oak Ridge National Laboratory. {title}. "
            "Oak Ridge (TN): Oak Ridge National Laboratory; 2024. "
            "Report No. ORNL/TM-2023/3218. "
            f"https://doi.org/{doi}"
        )
    return (
        f"{number}. {', '.join(names)}. {title}. {journal}. "
        f"{year};{locator}. https://doi.org/{doi}"
    )


def _datacite_reference(number: int, data: dict) -> str:
    attributes = data["attributes"]
    creators = attributes.get("creators", [])
    names = []
    for creator in creators[:6]:
        family = creator.get("familyName", "")
        if family and family.islower():
            family = family.title()
        given = creator.get("givenName", "")
        initials = "".join(part[0] for part in given.split() if part)
        names.append(f"{family} {initials}".strip() or creator.get("name", ""))
    if len(creators) > 6:
        names.append("et al")
    titles = attributes.get("titles", [])
    title = html.unescape(titles[0].get("title", "")) if titles else ""
    publisher = html.unescape(attributes.get("publisher", ""))
    year = attributes.get("publicationYear", "")
    doi = attributes.get("doi", data.get("id", ""))
    return (
        f"{number}. {', '.join(names)}. {title} [dataset]. "
        f"{publisher}; {year}. https://doi.org/{doi}"
    )


def _formatted_references(keys: list[str]) -> list[str]:
    snapshot = _latest_reference_snapshot()
    references = []
    for number, key in enumerate(keys, start=1):
        payload = json.loads((snapshot / f"{key}.json").read_text())
        if "message" in payload:
            references.append(_crossref_reference(number, payload["message"]))
        else:
            references.append(_datacite_reference(number, payload["data"]))
    return references


def _value_map() -> dict[str, str]:
    values = pd.read_csv(MANUSCRIPT_VALUES).fillna("")
    return dict(zip(values["value_id"], values["display_value"], strict=True))


def _resolve_source() -> tuple[list[tuple[str, str]], list[str], int]:
    keys = _reference_keys()
    citation_numbers = {key: number for number, key in enumerate(keys, start=1)}
    values = _value_map()
    seen: list[str] = []
    resolved_blocks: list[tuple[str, str]] = []
    paragraphs = re.split(r"\n\s*\n", MANUSCRIPT_SOURCE.read_text().strip())
    for block in paragraphs:
        block = block.strip()
        if not block:
            continue
        kind = "paragraph"
        if block.startswith("## "):
            kind = "heading2"
            block = block[3:]
        elif block.startswith("# "):
            kind = "heading1"
            block = block[2:]
        elif block.startswith("[[FIGURE:"):
            kind = "figure"
            block = block.removeprefix("[[FIGURE:").removesuffix("]]")
        elif block.startswith("[[TABLE:"):
            kind = "table"
            block = block.removeprefix("[[TABLE:").removesuffix("]]")

        def replace_value(match: re.Match[str]) -> str:
            value_id = match.group(1)
            if value_id not in values:
                raise ValueError(f"Unknown manuscript value: {value_id}")
            return values[value_id]

        def replace_citation(match: re.Match[str]) -> str:
            citation_keys = match.group(1).split(",")
            numbers = []
            for key in citation_keys:
                if key not in citation_numbers:
                    raise ValueError(f"Unknown JCP citation key: {key}")
                if key not in seen:
                    seen.append(key)
                numbers.append(citation_numbers[key])
            return "{" + ",".join(str(number) for number in numbers) + "}"

        block = re.sub(r"\{\{value:([^}]+)\}\}", replace_value, block)
        block = re.sub(r"\{\{cite:([^}]+)\}\}", replace_citation, block)
        unresolved = re.findall(r"\{\{[^}]+\}\}", block)
        if unresolved:
            raise ValueError(f"Unresolved source tokens: {unresolved}")
        resolved_blocks.append((kind, block))

    if seen != keys:
        raise ValueError(
            "References must be cited once in registry order before reuse. "
            f"Expected {keys}; observed {seen}."
        )
    countable = [
        text
        for kind, text in resolved_blocks
        if kind in {"paragraph", "heading1", "heading2"}
        and text not in {"Funding", "Declaration of competing interest"}
    ]
    word_count = len(re.findall(r"\b[\w'-]+\b", "\n".join(countable)))
    if not 6000 <= word_count <= 8000:
        raise ValueError(f"JCP manuscript word count outside 6000-8000: {word_count}")
    return resolved_blocks, _formatted_references(keys), word_count


def _main_tables() -> dict[int, tuple[pd.DataFrame, str]]:
    datasets = pd.read_csv(TABLES / "table_1_datasets.csv")
    assumptions = pd.read_csv(TABLES / "table_2_model_assumptions.csv")
    settings = assumptions.set_index("assumption")

    def setting(name: str) -> str:
        row = settings.loc[name]
        return f"{row['value']:g} {row['unit']}"

    climate = datasets[datasets["provider"] == "NOAA Physical Sciences Laboratory"]
    population = datasets[datasets["provider"] == "WorldPop"].iloc[0]
    boundaries = datasets[datasets["provider"] == "Natural Earth"].iloc[0]
    table1 = pd.DataFrame(
        [
            {
                "Domain": "Climate and masks",
                "Primary input or setting": (
                    f"{climate.iloc[0]['product']}; {climate.iloc[0]['version']}; "
                    f"{len(climate)} harmonized fields"
                ),
                "Role": "Meteorology, radiation, wind, land, snow, and sea ice",
            },
            {
                "Domain": "Population and labels",
                "Primary input or setting": (
                    f"{population['product']}; {population['version']}; "
                    f"{boundaries['product']} {boundaries['version']}"
                ),
                "Role": "Population burden and descriptive geography",
            },
            {
                "Domain": "Atmospheric response",
                "Primary input or setting": (
                    f"Mixing height {setting('Atmospheric mixing height')}; "
                    f"density {setting('Air density')}; "
                    f"heat capacity {setting('Air heat capacity')}"
                ),
                "Role": "Cell heat capacity and one-unit temperature response",
            },
            {
                "Domain": "Canonical search",
                "Primary input or setting": (
                    f"Transfer {setting('Canonical transferred energy')}; "
                    f"Humidex reference {setting('Humidex reference')}; "
                    f"curvature {setting('Burden curvature')}; "
                    f"distance {setting('Maximum transfer distance')}"
                ),
                "Role": "Prespecified global screening configuration",
            },
            {
                "Domain": "Finite capacity",
                "Primary input or setting": (
                    f"Heat-load cap {setting('Local heat-load cap')}; "
                    "temperature perturbation "
                    f"{setting('Maximum local temperature perturbation')}"
                ),
                "Role": "Local finite-transfer feasibility",
            },
        ]
    )

    failure = pd.read_csv(FAILURE_MATRIX)
    table2 = failure[
        [
            "receiving_class",
            "practical_default",
            "free_sink_failure_mode",
            "required_evidence",
        ]
    ].copy()
    table2.columns = [
        "Receiving class",
        "Practical default",
        "Why it is not a free sink",
        "Evidence required",
    ]

    structural = pd.read_csv(STRUCTURAL)
    spatial_values = structural["spatial_zero_penalty_net_benefit_per_gj"]
    temporal_ratios = structural["temporal_to_spatial_value_ratio"]
    table3 = pd.DataFrame(
        [
            {
                "Diagnostic": "Prespecified scenarios",
                "Coverage": (
                    f"{len(structural)} scenarios; "
                    f"{structural['climate_year'].nunique()} climate years; "
                    f"{structural['population_product'].nunique()} "
                    "population products; "
                    f"{structural['thermal_metric'].nunique()} metrics"
                ),
                "Result": (
                    f"{(spatial_values > 0).sum()}/{len(structural)} retained "
                    "positive constrained spatial value"
                ),
            },
            {
                "Diagnostic": "Constrained spatial value",
                "Coverage": "All structural scenarios",
                "Result": (
                    f"{spatial_values.min():.3f}-{spatial_values.max():.3f} "
                    "burden units/GJ"
                ),
            },
            {
                "Diagnostic": "Temporal/spatial ratio",
                "Coverage": "Matched source events",
                "Result": (f"{temporal_ratios.min():.3f}-{temporal_ratios.max():.3f}"),
            },
            {
                "Diagnostic": "Unrestricted pathological sink",
                "Coverage": "All structural scenarios",
                "Result": (
                    f"{structural['unrestricted_pathological_sinks_present'].sum()}/"
                    f"{len(structural)} scenarios"
                ),
            },
        ]
    )

    evidence = pd.read_csv(CASE_EVIDENCE)
    evidence_implications = {
        "FOUND_MEASURED": (
            "Measured source series, outage information, and receiving-demand "
            "summary bounds are available; no synchronized demand trace."
        ),
        "FOUND_DERIVABLE": (
            "Heat-pump COP and capacity support bounded sensitivity analysis."
        ),
        "FOUND_BOUND_ONLY": (
            "Capture, auxiliaries, route, cost, and site context remain "
            "insufficient for case-specific project claims."
        ),
        "NOT_FOUND": (
            "Hydraulics, storage design, lifecycle inventory, recovery-system "
            "availability, and new-discharge evidence remain open."
        ),
    }
    evidence_counts = evidence["status"].value_counts()
    table4 = pd.DataFrame(
        [
            {
                "Evidence class": status,
                "Items": int(evidence_counts.get(status, 0)),
                "Permitted interpretation": implication,
            }
            for status, implication in evidence_implications.items()
        ]
    )

    pathways = pd.read_csv(PATHWAYS)
    table5 = pathways[
        [
            "pathway_label",
            "assessment_status",
            "binding_constraints",
            "conclusion",
        ]
    ].copy()
    table5.columns = [
        "Pathway",
        "Eligibility",
        "Binding evidence",
        "Permitted conclusion",
    ]
    return {
        1: (
            table1,
            "Table 1. Principal public datasets and prespecified "
            "global-model assumptions.",
        ),
        2: (
            table2,
            "Table 2. Receiving classes, free-sink failure modes, "
            "and evidence requirements.",
        ),
        3: (
            table3,
            "Table 3. Prespecified structural-replication results.",
        ),
        4: (
            table4,
            "Table 4. Measured real-case evidence and claim limits.",
        ),
        5: (
            table5,
            "Table 5. Supported, conditional, and excluded receiving pathways.",
        ),
    }


def _add_title(document: Document, *, author_details: bool) -> None:
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run(TITLE)
    if author_details:
        author = document.add_paragraph()
        author.alignment = WD_ALIGN_PARAGRAPH.CENTER
        author.add_run("Tatsuki Onishi").bold = True
        affiliation = document.add_paragraph()
        affiliation.alignment = WD_ALIGN_PARAGRAPH.CENTER
        affiliation.add_run("AUTHOR ACTION: confirm affiliation.")
        corresponding = document.add_paragraph()
        corresponding.alignment = WD_ALIGN_PARAGRAPH.CENTER
        corresponding.add_run(
            "Corresponding author: Tatsuki Onishi; bougtoir@gmail.com"
        )
        document.add_page_break()


def _build_manuscript(
    blocks: list[tuple[str, str]],
    references: list[str],
    *,
    inline_assets: bool,
    word_count: int,
) -> Path:
    document = Document()
    _configure_document(document, compact=inline_assets)
    _add_title(document, author_details=inline_assets)
    tables = _main_tables()
    for kind, text in blocks:
        if kind == "heading1":
            document.add_heading(text, level=1)
        elif kind == "heading2":
            document.add_heading(text, level=2)
        elif kind == "paragraph":
            if text.startswith("Keywords:"):
                paragraph = document.add_paragraph()
                prefix = paragraph.add_run("Keywords:")
                prefix.bold = True
                paragraph.add_run(text.removeprefix("Keywords:"))
            else:
                _add_cited_paragraph(document, text)
        elif kind == "figure":
            number = int(text)
            if inline_assets:
                _add_inline_figure(document, number)
            else:
                paragraph = document.add_paragraph()
                paragraph.add_run(f"[Insert Figure {number} near here]").italic = True
        elif kind == "table":
            number = int(text)
            if inline_assets:
                _add_table(document, *tables[number])
            else:
                paragraph = document.add_paragraph()
                paragraph.add_run(f"[Insert Table {number} near here]").italic = True

    document.add_heading("References", level=1)
    for reference in references:
        document.add_paragraph(reference)

    if not inline_assets:
        document.add_page_break()
        document.add_heading("Tables", level=1)
        for number in sorted(tables):
            _add_table(document, *tables[number])
        document.add_heading("Figure legends", level=1)
        for number in sorted(FIGURE_CAPTIONS):
            document.add_paragraph(FIGURE_CAPTIONS[number])

    document.core_properties.comments = (
        f"Programmatically generated JCP manuscript; word count {word_count}."
    )
    output = SUBMISSION / (
        "manuscript_jcp_inline.docx" if inline_assets else "manuscript_jcp.docx"
    )
    document.save(output)
    return output


def _build_title_page(word_count: int) -> Path:
    document = Document()
    _configure_document(document)
    _add_title(document, author_details=True)
    document.add_heading("Article type", level=1)
    document.add_paragraph("Original article")
    document.add_heading("Manuscript information", level=1)
    document.add_paragraph(f"Main-text word count: {word_count}")
    document.add_paragraph("Number of main figures: 6")
    document.add_paragraph("Number of main tables: 5")
    document.add_paragraph("Number of references: 48")
    document.add_heading("Author-controlled metadata", level=1)
    document.add_paragraph(
        "AUTHOR ACTION: confirm affiliation, postal address, ORCID, funding, "
        "competing interests, and permanent repository DOI."
    )
    output = SUBMISSION / "title_page.docx"
    document.save(output)
    return output


def _build_highlights() -> Path:
    lines = [
        "Highlights",
        "",
        "• Conserved-heat optimization exposes apparent free-sink optima",
        "• Low human exposure does not establish environmental eligibility",
        "• Temporal delay retains 53.9% of matched spatial screening value",
        "• Evidence gates leave zero supported practical endpoints",
        "• Local heat removal does not imply environmental burden removal",
    ]
    output = SUBMISSION / "highlights.txt"
    output.write_text("\n".join(lines) + "\n")
    return output


def _build_cover_letter() -> Path:
    document = Document()
    _configure_document(document)
    document.add_paragraph("Dear Editors of the Journal of Cleaner Production,")
    document.add_paragraph(
        f"Please consider our Original Article, “{TITLE}”. The manuscript "
        "addresses cleaner production as a system-boundary problem: removing "
        "heat locally can shift burden to an unevaluated receiving environment."
    )
    document.add_paragraph(
        "The contribution is a conservation-aware, evidence-gated framework "
        "that separates objective value from destination eligibility, preserves "
        "incompatible burden, energy, cost, and emissions units, and connects "
        "a global human-burden screening model to a measured waste-heat "
        "pathway. Heat is a useful test because conservation makes transfer "
        "explicit. A problem does not disappear when it crosses the system "
        "boundary. The framework exposes an unconstrained open-ocean optimum "
        "as a diagnostic failure and reports zero supported practical endpoints "
        "rather than filling missing engineering or lifecycle evidence."
    )
    document.add_paragraph(
        "The study does not claim a safe global sink, deployable long-distance "
        "heat relocation, planetary cooling, health benefit, case-specific "
        "economics, quantified ecosystem damage, or case-specific lifecycle "
        "performance. Its relevance to JCP is the prevention and transparent "
        "detection of problem shifting in waste-heat management, industrial "
        "ecology, and sustainability claims."
    )
    document.add_paragraph(
        "AUTHOR ACTION: confirm originality, author approval, funding, "
        "competing interests, and any required suggested reviewers before "
        "submission."
    )
    document.add_paragraph("Sincerely,\nTatsuki Onishi")
    output = SUBMISSION / "cover_letter.docx"
    document.save(output)
    return output


def _add_pptx_box(
    slide,
    *,
    left: float,
    top: float,
    width: float,
    height: float,
    title: str,
    body: str,
    fill: str,
) -> None:
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        PptxInches(left),
        PptxInches(top),
        PptxInches(width),
        PptxInches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = PptxRGBColor.from_string(fill)
    shape.line.color.rgb = PptxRGBColor.from_string("374151")
    frame = shape.text_frame
    frame.clear()
    first = frame.paragraphs[0]
    first.text = title
    first.alignment = PP_ALIGN.CENTER
    first.runs[0].font.size = PptxPt(17)
    first.runs[0].font.bold = True
    second = frame.add_paragraph()
    second.text = body
    second.alignment = PP_ALIGN.CENTER
    second.runs[0].font.size = PptxPt(12)


def _build_graphical_abstract() -> tuple[Path, Path]:
    boxes = [
        ("HEAT SOURCE", "Measured or modeled\nconserved heat", "DBEAFE"),
        ("LOCAL RELIEF", "Minimize local\nhuman burden", "EDE9FE"),
        (
            "APPARENT SINK",
            "Remote land | ocean\ncryosphere | atmosphere\nradiative concept",
            "FEF3C7",
        ),
        (
            "NOT A FREE SINK",
            "Demand | route | capture\nauxiliaries | ecology\ncost | lifecycle",
            "FEE2E2",
        ),
        (
            "VALIDATED PATHWAY",
            "Reuse or storage\nwhen evidenced\n0 evidence-supported modes",
            "DCFCE7",
        ),
    ]
    presentation = Presentation()
    presentation.slide_width = PptxInches(13.333)
    presentation.slide_height = PptxInches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    title = slide.shapes.add_textbox(
        PptxInches(0.5), PptxInches(0.25), PptxInches(12.3), PptxInches(0.7)
    )
    paragraph = title.text_frame.paragraphs[0]
    paragraph.text = "Conservation-aware detection of thermal burden shifting"
    paragraph.alignment = PP_ALIGN.CENTER
    paragraph.runs[0].font.size = PptxPt(25)
    paragraph.runs[0].font.bold = True
    lefts = [0.35, 2.95, 5.55, 8.15, 10.75]
    for left, (heading, body, fill) in zip(lefts, boxes, strict=True):
        _add_pptx_box(
            slide,
            left=left,
            top=2.0,
            width=2.25,
            height=2.3,
            title=heading,
            body=body,
            fill=fill,
        )
    for start, end in zip(lefts[:-1], lefts[1:], strict=True):
        connector = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            PptxInches(start + 2.25),
            PptxInches(3.15),
            PptxInches(end),
            PptxInches(3.15),
        )
        connector.line.color.rgb = PptxRGBColor.from_string("4B5563")
        connector.line.width = PptxPt(2)
    footer = slide.shapes.add_textbox(
        PptxInches(0.6), PptxInches(5.55), PptxInches(12.1), PptxInches(0.8)
    )
    footer_paragraph = footer.text_frame.paragraphs[0]
    footer_paragraph.text = (
        "Removing heat from one place does not remove its environmental consequences."
    )
    footer_paragraph.alignment = PP_ALIGN.CENTER
    footer_paragraph.runs[0].font.size = PptxPt(19)
    footer_paragraph.runs[0].font.bold = True
    pptx_output = SUBMISSION / "graphical_abstract_editable.pptx"
    presentation.save(pptx_output)

    figure, axis = plt.subplots(figsize=(13.333, 7.5))
    axis.set_axis_off()
    axis.text(
        0.5,
        0.91,
        "Conservation-aware detection of thermal burden shifting",
        ha="center",
        transform=axis.transAxes,
        fontsize=21,
        fontweight="bold",
    )
    centers = [0.10, 0.30, 0.50, 0.70, 0.90]
    for center, (heading, body, fill) in zip(centers, boxes, strict=True):
        axis.add_patch(
            plt.Rectangle(
                (center - 0.085, 0.43),
                0.17,
                0.26,
                transform=axis.transAxes,
                facecolor=f"#{fill}",
                edgecolor="#374151",
                linewidth=1.2,
            )
        )
        axis.text(
            center,
            0.64,
            heading,
            ha="center",
            va="center",
            transform=axis.transAxes,
            fontsize=10,
            fontweight="bold",
        )
        axis.text(
            center,
            0.535,
            body,
            ha="center",
            va="center",
            transform=axis.transAxes,
            fontsize=8.0,
        )
    for start, end in zip(centers[:-1], centers[1:], strict=True):
        axis.annotate(
            "",
            xy=(end - 0.087, 0.56),
            xytext=(start + 0.087, 0.56),
            xycoords=axis.transAxes,
            arrowprops={"arrowstyle": "->", "lw": 1.8, "color": "#4b5563"},
        )
    axis.text(
        0.5,
        0.12,
        "Removing heat from one place does not remove its environmental consequences.",
        ha="center",
        transform=axis.transAxes,
        fontsize=17,
        fontweight="bold",
    )
    png_output = SUBMISSION / "graphical_abstract.png"
    figure.savefig(png_output, dpi=300, bbox_inches="tight", facecolor="white")
    plt.close(figure)
    return pptx_output, png_output


def _build_supplement() -> Path:
    document = Document()
    _configure_document(document)
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    document.add_heading(f"Supplementary material: {TITLE}", level=0)
    document.add_paragraph(
        "The supplement reports diagnostics generated from the frozen "
        "analysis and introduces no additional empirical observations."
    )
    document.add_heading("Supplementary methods", level=1)
    document.add_paragraph(
        "The exact spatial search processes each six-hour field in time order. "
        "A spherical nearest-neighbour structure identifies candidates within "
        "the maximum distance, after which exact great-circle distance is "
        "evaluated. Temporal candidates are vectorized over lag steps."
    )
    document.add_paragraph(
        "Finite-Q burden changes are evaluated analytically under the "
        "quadratic burden function. Validation covers objective identities, "
        "distance, same-location temporal behavior, energy conservation, "
        "no-relocation retention, finite-Q feasibility, seasonal groups, "
        "trajectory conservation, and noncausal radiation interpretation."
    )
    document.add_heading("Supplementary tables", level=1)
    datasets = pd.read_csv(TABLES / "table_1_datasets.csv")
    assumptions = pd.read_csv(TABLES / "table_2_model_assumptions.csv")
    inventory = pd.concat(
        [
            pd.DataFrame(
                {
                    "domain": "dataset",
                    "item": datasets["product"],
                    "setting_or_role": datasets["role"],
                    "version_or_unit": datasets["version"],
                    "status": datasets["completion"],
                }
            ),
            pd.DataFrame(
                {
                    "domain": "model assumption",
                    "item": assumptions["assumption"],
                    "setting_or_role": assumptions["value"],
                    "version_or_unit": assumptions["unit"],
                    "status": "prespecified",
                }
            ),
        ],
        ignore_index=True,
    )
    evidence = pd.read_csv(CASE_EVIDENCE)
    structural = pd.read_csv(STRUCTURAL)
    dispatch_validation = pd.read_csv(
        RESULTS / "diagnostics" / "finite_dispatch_validation.csv"
    )
    supplement_tables = [
        (
            inventory,
            "Table S1. Full dataset inventory and prespecified model settings.",
        ),
        (
            evidence[
                [
                    "required_item",
                    "status",
                    "value_summary",
                    "claim_limit",
                    "downstream_decision",
                ]
            ],
            "Table S2. Full item-level real-case evidence registry.",
        ),
        (
            structural[
                [
                    "scenario_id",
                    "climate_year",
                    "population_product",
                    "spatial_coarsening_factor",
                    "thermal_metric",
                    "spatial_zero_penalty_net_benefit_per_gj",
                    "temporal_to_spatial_value_ratio",
                    "unrestricted_pathological_sinks_present",
                ]
            ],
            "Table S3. Full prespecified structural-replication results.",
        ),
        (
            dispatch_validation,
            "Table S4. Finite-dispatch validation checks.",
        ),
        (
            pd.read_csv(TABLES / "table_6_robustness_summary.csv"),
            "Table S5. Parametric source-hotspot robustness.",
        ),
        (
            pd.read_csv(RESULTS / "canonical" / "safety_ablation.csv")[
                [
                    "scenario",
                    "sink_latitude",
                    "sink_longitude",
                    "sink_is_ocean",
                    "sink_is_cryosphere",
                    "distance_km",
                    "source_marginal_benefit",
                    "sink_marginal_burden",
                    "net_benefit",
                ]
            ],
            "Table S6. Progressive receiving-mask ablation.",
        ),
        (
            pd.read_csv(RESULTS / "canonical" / "finite_q_sensitivity.csv")[
                [
                    "energy_gj",
                    "status",
                    "distance_km",
                    "source_temperature_change_k",
                    "sink_temperature_change_k",
                    "source_benefit_total",
                    "sink_burden_total",
                    "transport_penalty_total",
                    "net_benefit_per_gj",
                ]
            ],
            "Table S7. Finite-transfer and local-capacity sensitivity.",
        ),
        (
            pd.read_csv(TABLES / "burden_shifting_cascade.csv")[
                [
                    "stage_order",
                    "stage_label",
                    "receiving_class",
                    "screening_value_burden_units_per_gj",
                    "value_retained_fraction",
                    "objective_detects_the_shift",
                    "evidence_status",
                    "eligibility",
                ]
            ],
            "Table S8. Machine-readable burden-shifting cascade.",
        ),
        (
            pd.read_csv(PRACTICAL_MODES)[
                [
                    "mode_label",
                    "demand_basis",
                    "storage_basis",
                    "eligibility",
                    "evidence_gap_count",
                    "useful_heat_delivered_mwh",
                    "demand_served_fraction",
                    "residual_source_rejection_mwh",
                    "mode_decision",
                ]
            ],
            "Table S9. Conditional practical-mode sensitivities.",
        ),
        (
            pd.read_csv(WATERFALL)[
                [
                    "scenario_id",
                    "stage_order",
                    "stage_label",
                    "value",
                    "unit",
                    "numeric_status",
                    "evidence_status",
                    "eligibility",
                ]
            ],
            "Table S10. Global-to-real constraint waterfall. The repository "
            "CSV retains the complete interpretation field.",
        ),
    ]
    for frame, caption in supplement_tables:
        _add_compact_table(document, frame, caption)
    document.add_heading("Supplementary figure legends", level=1)
    legends = [
        "Figure S1. Parametric source-hotspot robustness.",
        "Figure S2. Finite-transfer sensitivity under the local cap.",
        "Figure S3. Reduced-order wind-trajectory diagnostic.",
        "Figure S4. Seasonal and local day-night marginal burden.",
        "Figure S5. Progressive mask-ablation geography.",
        "Figure S6. Abstract joint transport-storage penalty sensitivity.",
    ]
    for legend in legends:
        document.add_paragraph(legend)
    output = SUBMISSION / "supplementary_material.docx"
    document.save(output)
    return output


def _copy_figures() -> list[Path]:
    FIGURE_DIR.mkdir(parents=True, exist_ok=True)
    outputs = []
    for number, stem in FIGURE_SOURCES.items():
        for suffix in (".png", ".svg", ".pdf"):
            source = FIGURES / f"{stem}{suffix}"
            output = FIGURE_DIR / f"figure_{number}{suffix}"
            shutil.copy2(source, output)
            outputs.append(output)
    supplementary = [
        "supplementary_figure_s1_source_hotspot_robustness",
        "supplementary_figure_s2_finite_q_sensitivity",
        "supplementary_figure_s3_wind_trajectory",
        "supplementary_figure_s5_seasonal_day_night",
        "supplementary_figure_s6_safety_ablation",
        "supplementary_figure_s7_joint_penalty_sensitivity",
    ]
    for number, stem in enumerate(supplementary, start=1):
        for suffix in (".png", ".svg", ".pdf"):
            source = FIGURES / f"{stem}{suffix}"
            output = FIGURE_DIR / f"supplementary_figure_s{number}{suffix}"
            shutil.copy2(source, output)
            outputs.append(output)
    return outputs


def _build_checklist(word_count: int) -> Path:
    content = f"""# Journal of Cleaner Production submission checklist

- [x] Original Article selected
- [x] Main-text word count is {word_count}, within the 6000-8000 target
- [x] Forty-eight verified references, below the 50-reference limit
- [x] Abstract and keywords included
- [x] Highlights supplied separately
- [x] Graphical abstract supplied as PNG and editable PPTX
- [x] Six main figures supplied as PNG, SVG, and vector PDF
- [x] Five main tables included in the manuscript
- [x] Tables avoid vertical rules and cell shading
- [x] Supplementary diagnostics separated from the main figures
- [x] Data and code availability statement included
- [x] Generative-AI disclosure included
- [x] No burden-unit/MWh-th/USD/CO2e composite score
- [x] No claim of planetary cooling, health benefit, or a validated safe sink
- [x] Final live JCP Guide for Authors check completed on 2026-09-25
- [ ] AUTHOR ACTION: confirm affiliation and postal address
- [ ] AUTHOR ACTION: confirm the CRediT authorship contribution statement
- [ ] AUTHOR ACTION: confirm funding statement
- [ ] AUTHOR ACTION: confirm competing-interest declaration
- [ ] AUTHOR ACTION: archive the repository and insert its DOI
"""
    output = SUBMISSION / "submission_checklist.md"
    output.write_text(content)
    return output


def _build_archive(files: list[Path]) -> Path:
    output = SUBMISSION / "jcp_submission_package.zip"
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for path in sorted(files):
            archive.write(path, path.relative_to(SUBMISSION))
    _normalize_zip_archive(output)
    return output


def main() -> None:
    SUBMISSION.mkdir(parents=True, exist_ok=True)
    blocks, references, word_count = _resolve_source()
    outputs = [
        _build_manuscript(
            blocks,
            references,
            inline_assets=False,
            word_count=word_count,
        ),
        _build_manuscript(
            blocks,
            references,
            inline_assets=True,
            word_count=word_count,
        ),
        _build_title_page(word_count),
        _build_highlights(),
        _build_cover_letter(),
        *_build_graphical_abstract(),
        _build_supplement(),
        _build_checklist(word_count),
        *_copy_figures(),
    ]
    for output in outputs:
        if output.suffix in {".docx", ".pptx"}:
            _normalize_zip_archive(output)
    archive = _build_archive(outputs)
    print(f"JCP manuscript word count: {word_count}")
    for output in [*outputs, archive]:
        print(f"Wrote {output.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
