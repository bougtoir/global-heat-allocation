"""Build the Applied Energy manuscript and submission package."""

from __future__ import annotations

import argparse
import json
import re
import shutil
import zipfile
from pathlib import Path

import matplotlib.pyplot as plt
import pandas as pd
import xarray as xr
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor as PptxRGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches as PptxInches
from pptx.util import Pt as PptxPt

ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT / "results"
CANONICAL = RESULTS / "canonical"
TABLES = RESULTS / "tables"
FIGURES = RESULTS / "figures"
TECHNO_ECONOMIC = RESULTS / "techno_economic"
ENVIRONMENTAL_CONSTRAINTS = RESULTS / "environmental_constraints"
DISPATCH = RESULTS / "dispatch"
STRUCTURAL_REPLICATION = TABLES / "structural_replication.csv"
SUBMISSION = ROOT / "submission"
REFERENCE_ROOT = ROOT / "data" / "raw" / "references"
TECHNO_ECONOMIC_PARAMETERS = (
    ROOT / "data" / "metadata" / "techno_economic_parameters.csv"
)
ENVIRONMENTAL_PARAMETERS = ROOT / "data" / "metadata" / "environmental_parameters.csv"
MANUSCRIPT_VALUES = ROOT / "provenance" / "manuscript_values.csv"
CASE_EVIDENCE = ROOT / "provenance" / "case_specific_evidence.csv"
CONSTRAINT_WATERFALL = TABLES / "global_to_real_constraint_waterfall.csv"
PRACTICAL_PARETO = TECHNO_ECONOMIC / "practical_pareto_modes.csv"

REFERENCE_KEYS = [
    "flanner2009anthropogenic",
    "chen2004anthropogenic",
    "li2022allocation",
    "guelpa2019storage",
    "lund2014district",
    "peyre2019transport",
    "john2022oceanmover",
    "parno2019seaice",
    "vissio2020wasserstein",
    "buzan2015heatstress",
    "dinapoli2021era5heat",
    "raymond2020heat",
    "mora2017deadlyheat",
    "kalnay1996ncep",
    "tatem2017worldpop",
    "stull2011wetbulb",
    "loeb2018ceres",
    "raman2014radiative",
]


def _normalize_zip_archive(path: Path) -> None:
    with zipfile.ZipFile(path) as archive:
        members = [(item.filename, archive.read(item)) for item in archive.infolist()]
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for filename, data in members:
            item = zipfile.ZipInfo(filename, date_time=(1980, 1, 1, 0, 0, 0))
            item.compress_type = zipfile.ZIP_DEFLATED
            item.create_system = 3
            item.external_attr = 0o600 << 16
            archive.writestr(item, data)


def _latest_reference_snapshot() -> Path:
    snapshots = sorted(REFERENCE_ROOT.glob("crossref_*"))
    if not snapshots:
        raise FileNotFoundError("No Crossref snapshot is available.")
    return snapshots[-1]


def _reference_metadata() -> dict[str, dict]:
    snapshot = _latest_reference_snapshot()
    return {
        key: json.loads((snapshot / f"{key}.json").read_text())["message"]
        for key in REFERENCE_KEYS
    }


def _format_reference(number: int, metadata: dict) -> str:
    authors = metadata.get("author", [])
    names = []
    for author in authors[:6]:
        family = author.get("family", "")
        initials = "".join(part[0] for part in author.get("given", "").split() if part)
        names.append(f"{family} {initials}".strip())
    if len(authors) > 6:
        names.append("et al.")
    title = metadata.get("title", [""])[0].rstrip(".")
    journal = metadata.get("container-title", [""])[0]
    issued = metadata.get("issued", {}).get("date-parts", [[None]])[0]
    year = issued[0]
    volume = metadata.get("volume", "")
    issue = metadata.get("issue", "")
    pages = metadata.get("page", "")
    volume_issue = volume + (f"({issue})" if issue else "")
    locator = f"{volume_issue}:{pages}" if pages else volume_issue
    doi = metadata.get("DOI", "")
    return (
        f"{number}. {', '.join(names)}. {title}. {journal}. "
        f"{year};{locator}. https://doi.org/{doi}"
    )


def _set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _add_cited_paragraph(
    document: Document,
    text: str,
    *,
    style: str | None = None,
    bold_prefix: str | None = None,
) -> None:
    paragraph = document.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        prefix_run = paragraph.add_run(bold_prefix)
        prefix_run.bold = True
        text = text[len(bold_prefix) :]
    for part in re.split(r"(\{[^}]+\})", text):
        if part.startswith("{") and part.endswith("}"):
            run = paragraph.add_run(part[1:-1])
            run.font.superscript = True
        else:
            paragraph.add_run(part)


def _configure_document(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 2
    normal.paragraph_format.space_after = Pt(0)
    for heading in ["Title", "Heading 1", "Heading 2", "Heading 3"]:
        styles[heading].font.name = "Arial"
        styles[heading].font.color.rgb = RGBColor(0, 0, 0)
    for section in document.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)


def _add_table(document: Document, frame: pd.DataFrame, caption: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run(caption).bold = True
    table = document.add_table(rows=1, cols=len(frame.columns))
    table.style = "Table Grid"
    for column, cell in zip(frame.columns, table.rows[0].cells, strict=True):
        cell.text = str(column).replace("_", " ").title()
        _set_cell_shading(cell, "D9EAF2")
    header_properties = table.rows[0]._tr.get_or_add_trPr()
    repeat_header = OxmlElement("w:tblHeader")
    repeat_header.set(qn("w:val"), "true")
    header_properties.append(repeat_header)
    for row in frame.itertuples(index=False):
        cells = table.add_row().cells
        for value, cell in zip(row, cells, strict=True):
            if isinstance(value, float):
                cell.text = f"{value:.3g}"
            else:
                cell.text = str(value)
    for row in table.rows:
        row_properties = row._tr.get_or_add_trPr()
        row_properties.append(OxmlElement("w:cantSplit"))


def _figure_legends(context: dict) -> list[str]:
    return [
        (
            "Figure 1. Two-scale heat-allocation and evidence framework. The "
            "global burden-space screening branch and the real physical-energy "
            "branch enter a common evidence-gated constraint waterfall without "
            "an assumed unit conversion."
        ),
        (
            "Figure 2. Primary spatial one-GJ transfer under a small transport "
            "penalty of 10^-5 burden units per GJ-km, a 5,000-km limit, and "
            "land/non-cryosphere sink constraints."
        ),
        (
            "Figure 3. Net marginal burden reduction for zero-penalty spatial, "
            "temporal, and joint modes using the same peak source cell-time; "
            f"the temporal candidate uses a {context['temporal_lag_hours']}-hour lag."
        ),
        (
            "Figure 4. Measured Frontier waste-heat power and reported ORNL "
            "receiving-demand bounds. Monthly source summaries retain "
            "unavailable 10-minute intervals rather than imputing them; "
            "horizontal lines are annual demand-summary bounds, not a "
            "synchronized demand profile."
        ),
        (
            "Figure 5. Global-to-real constraint waterfall. Burden-space and "
            "physical-energy branches remain separate because no supported "
            "unit bridge exists. The real pathway retains conditional inputs "
            "but reaches no supported practical endpoint."
        ),
        (
            "Figure 6. Conditional practical mode results under reported "
            "lower and upper demand bounds. All modes are sensitivity cases; "
            "storage labels do not imply a case-specific installed design."
        ),
    ]


def _main_tables(context: dict) -> list[tuple[pd.DataFrame, str]]:
    datasets = pd.read_csv(TABLES / "table_1_datasets.csv")
    assumptions = pd.read_csv(TABLES / "table_2_model_assumptions.csv")
    comparison = pd.read_csv(TABLES / "table_4_allocation_comparison.csv")
    values = context["manuscript_values"]

    def value(value_id: str) -> str:
        return values.loc[value_id, "display_value"]

    case_summary = pd.DataFrame(
        [
            [
                "Source heat",
                "Measured",
                f"{value('frontier_observed_heat')} MWh-th; "
                f"{value('frontier_valid_intervals')} valid intervals",
            ],
            [
                "Receiving demand",
                "Summary bounds only",
                f"Case A {value('ornl_case_a_average')}–"
                f"{value('ornl_case_a_maximum')} MW; no synchronized profile",
            ],
            [
                "Heat pump",
                "Bounded",
                f"COP {value('heat_pump_case_cop')}; "
                f"{value('heat_pump_case_total_capacity')} MW-th",
            ],
            [
                "Capture, route, auxiliaries",
                "Not calibrated",
                "No recovery efficiency, hydraulic design, or incremental load",
            ],
            [
                "Storage",
                "Not designed",
                "No case-specific tank volume or operating temperature range",
            ],
            [
                "Cost and lifecycle CO2e",
                "Incomplete",
                "No complete installed cost or case-specific lifecycle inventory",
            ],
            [
                "Environmental clearance",
                "Not established",
                "No route clearance or authorization for a new thermal discharge",
            ],
        ],
        columns=["Evidence domain", "Status", "Supported statement"],
    )
    waterfall_summary = pd.DataFrame(
        [
            [
                "Global spatial bound",
                f"{value('matched_spatial_value')} burden units/GJ",
                "Theoretical only",
                "No conversion to delivered heat",
            ],
            [
                "Global temporal bound",
                f"{value('matched_temporal_value')} burden units/GJ",
                "Theoretical only",
                "Same-event marginal screening",
            ],
            [
                "Observed Frontier source",
                f"{value('frontier_observed_heat')} MWh-th",
                "Supported input",
                "Missing intervals retained",
            ],
            [
                "ORNL demand",
                f"{value('ornl_case_a_average')}–"
                f"{value('ornl_case_b_selected_maximum')} MW",
                "Summary bounds",
                "No synchronized demand trace",
            ],
            [
                "Conditional finite dispatch",
                f"{value('dispatch_served_fraction_min')}–"
                f"{value('dispatch_served_fraction_max')} demand served",
                "Conditional",
                "Six sensitivity modes",
            ],
            [
                "Supported practical endpoint",
                value("supported_practical_mode_count"),
                "Excluded",
                "Required evidence gates remain open",
            ],
        ],
        columns=["Quantity", "Value", "Eligibility", "Claim limit"],
    )
    pareto_summary = context["pareto"].copy()
    pareto_summary["Mode"] = pareto_summary["storage_basis"].map(
        {
            "none": "Direct",
            "minimum_formula": "Formula storage",
            "generic_reference": "Generic storage",
        }
    )
    pareto_summary["Demand"] = pareto_summary["demand_basis"].map(
        {
            "reported_lower_bound": "Lower bound",
            "reported_upper_bound": "Upper bound",
        }
    )
    pareto_summary["Useful heat (MWh)"] = pareto_summary[
        "useful_heat_delivered_mwh"
    ].map(lambda item: f"{item:,.1f}")
    pareto_summary["Demand served"] = pareto_summary["demand_served_fraction"].map(
        lambda item: f"{item:.1%}"
    )
    pareto_summary["Residual heat (MWh)"] = pareto_summary[
        "residual_source_rejection_mwh"
    ].map(lambda item: f"{item:,.1f}")
    pareto_summary["Eligibility"] = "Conditional"
    return [
        (
            datasets[["product", "role", "version", "completion"]],
            "Table 1. Public datasets used in the global analysis.",
        ),
        (assumptions, "Table 2. Canonical model assumptions."),
        (
            comparison[
                [
                    "analysis",
                    "source_country",
                    "sink_country",
                    "lag_steps",
                    "distance_km",
                    "net_benefit",
                    "fraction_of_matched_spatial",
                    "cooptimal_candidate_count",
                    "near_optimal_candidate_count",
                ]
            ],
            "Table 3. Matched-event spatial, temporal, and joint allocation results.",
        ),
        (
            case_summary,
            "Table 4. Case-specific evidence status.",
        ),
        (
            waterfall_summary,
            "Table 5. Global-to-real constraint summary.",
        ),
        (
            pareto_summary[
                [
                    "Mode",
                    "Demand",
                    "Useful heat (MWh)",
                    "Demand served",
                    "Residual heat (MWh)",
                    "Eligibility",
                ]
            ],
            "Table 6. Conditional practical mode-decision results.",
        ),
    ]


def _add_inline_figure(
    document: Document,
    filename: str,
    caption: str,
) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.add_run().add_picture(str(FIGURES / filename), width=Inches(6.2))
    caption_paragraph = document.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.space_before = Pt(12)
    caption_paragraph.add_run(caption).bold = True


def _result_context() -> dict:
    allocations = pd.read_csv(CANONICAL / "one_unit_allocations.csv")
    robustness = pd.read_csv(CANONICAL / "source_hotspot_robustness.csv")
    safety = pd.read_csv(CANONICAL / "safety_ablation.csv")
    finite_q = pd.read_csv(CANONICAL / "finite_q_sensitivity.csv")
    seasonal = pd.read_csv(CANONICAL / "seasonal_day_night_summary.csv")
    redistribution = pd.read_csv(CANONICAL / "natural_redistribution_summary.csv")
    matched = pd.read_csv(CANONICAL / "matched_event_allocations.csv").set_index(
        "analysis"
    )
    structural = pd.read_csv(STRUCTURAL_REPLICATION)
    manuscript_values = pd.read_csv(MANUSCRIPT_VALUES).set_index("value_id")
    case_evidence = pd.read_csv(CASE_EVIDENCE)
    waterfall = pd.read_csv(CONSTRAINT_WATERFALL)
    pareto = pd.read_csv(PRACTICAL_PARETO)
    spatial = allocations[
        (allocations["analysis"] == "spatial")
        & (allocations["scenario"] == "primary_land_noncryosphere")
        & (allocations["transport_penalty_per_gj_km"] == 0.0)
    ].iloc[0]
    spatial_small_penalty = allocations[
        (allocations["analysis"] == "spatial")
        & (allocations["scenario"] == "primary_land_noncryosphere")
        & (allocations["transport_penalty_per_gj_km"] == 1.0e-5)
    ].iloc[0]
    temporal = allocations[
        (allocations["analysis"] == "temporal")
        & (allocations["maximum_lag_steps"] == 3)
        & (allocations["storage_penalty_per_gj_step"] == 0.0)
    ].iloc[0]
    joint = allocations[
        (allocations["analysis"] == "joint")
        & (allocations["transport_penalty_per_gj_km"] == 0.0)
        & (allocations["storage_penalty_per_gj_step"] == 0.0)
    ].iloc[0]
    costly_spatial = allocations[
        (allocations["analysis"] == "spatial")
        & (allocations["transport_penalty_per_gj_km"] == 5.0)
    ].iloc[0]
    costly_joint = allocations[
        (allocations["analysis"] == "joint")
        & (allocations["transport_penalty_per_gj_km"] == 5.0)
        & (allocations["storage_penalty_per_gj_step"] == 0.0)
    ].iloc[0]
    with xr.open_dataset(ROOT / "data" / "processed" / "analysis_cube_2023.nc") as cube:
        population = float(cube["population"].sum())
        latitude_count = cube.sizes["lat"]
        longitude_count = cube.sizes["lon"]
        time_count = cube.sizes["time"]
    temporal_lag_hours = (
        pd.Timestamp(matched.loc["temporal", "sink_time"])
        - pd.Timestamp(matched.loc["temporal", "source_time"])
    ) / pd.Timedelta(hours=1)
    return {
        "allocations": allocations,
        "robustness": robustness,
        "safety": safety,
        "finite_q": finite_q,
        "seasonal": seasonal,
        "redistribution": redistribution,
        "structural": structural,
        "manuscript_values": manuscript_values,
        "case_evidence": case_evidence,
        "waterfall": waterfall,
        "pareto": pareto,
        "matched": matched,
        "spatial": spatial,
        "spatial_small_penalty": spatial_small_penalty,
        "temporal": temporal,
        "joint": joint,
        "costly_spatial": costly_spatial,
        "costly_joint": costly_joint,
        "matched_spatial": matched.loc["spatial"],
        "matched_temporal": matched.loc["temporal"],
        "matched_joint": matched.loc["joint"],
        "temporal_fraction": matched.loc["temporal", "fraction_of_matched_spatial"],
        "temporal_lag_hours": int(temporal_lag_hours),
        "population": population,
        "latitude_count": latitude_count,
        "longitude_count": longitude_count,
        "time_count": time_count,
    }


def _build_manuscript(
    context: dict,
    references: dict[str, dict],
    *,
    inline_assets: bool = False,
) -> Path:
    document = Document()
    _configure_document(document)
    warning = document.add_paragraph()
    warning.add_run(
        "DRAFT—NOT FOR SUBMISSION UNTIL THE SCIENTIFIC AND ADMINISTRATIVE GATES "
        "IN THE CHECKLIST ARE RESOLVED"
    ).bold = True
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run(
        "From global heat-allocation bounds to a real waste-heat pathway: "
        "a two-scale constraint analysis"
    )
    author = document.add_paragraph()
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author.add_run("Tatsuki Onishi").bold = True
    affiliation = document.add_paragraph()
    affiliation.alignment = WD_ALIGN_PARAGRAPH.CENTER
    affiliation.add_run("[Affiliation to be confirmed before submission]")
    corresponding = document.add_paragraph()
    corresponding.alignment = WD_ALIGN_PARAGRAPH.CENTER
    corresponding.add_run("Corresponding author: Tatsuki Onishi; bougtoir@gmail.com")
    document.add_page_break()

    document.add_heading("Abstract", level=1)
    spatial = context["spatial"]
    small_penalty = context["spatial_small_penalty"]
    robustness = context["robustness"]
    finite_q = context["finite_q"]
    structural = context["structural"]
    main_tables = _main_tables(context)
    legends = _figure_legends(context)
    cooptimal_count = int(context["matched_spatial"]["cooptimal_candidate_count"])
    infeasible = finite_q[finite_q["status"] != "feasible"]
    values = context["manuscript_values"]
    observed_heat = values.loc["frontier_observed_heat", "display_value"]
    valid_intervals = values.loc["frontier_valid_intervals", "display_value"]
    missing_intervals = values.loc["frontier_missing_intervals", "display_value"]
    abstract = (
        "The value of moving or storing heat depends on where and when it is "
        "released, but a theoretical allocation optimum need not survive "
        "engineering and environmental constraints. We connect two scales: "
        "a global conserved-heat marginal screening model and an auditable "
        "waste-heat case based on measured Frontier supercomputer data and "
        "documented Oak Ridge National Laboratory heating-demand bounds. The "
        "global model combines six-hourly reanalysis, gridded population, a "
        "convex Humidex burden, exact memory-bounded spatial search, and "
        "same-location temporal storage. The matched-event land and "
        "non-cryosphere spatial bound was "
        f"{context['matched_spatial']['net_benefit']:.1f} burden units/GJ; "
        f"an {context['temporal_lag_hours']}-hour delay retained "
        f"{context['temporal_fraction']:.1%} of that value. The spatial "
        f"solution contained {cooptimal_count} exactly co-optimal sink cells, "
        "and the largest tested transfer was infeasible under the local "
        "temperature cap. The real case contained "
        f"{observed_heat} MWh-th across {valid_intervals} valid 10-minute "
        f"intervals, with {missing_intervals} unavailable intervals. We then "
        "applied environmental, geographic, pathway, technology, storage, "
        "demand, and finite-dispatch gates. Receiving demand was available "
        "only as measured summary bounds; route hydraulics, capture "
        "efficiency, incremental auxiliaries, storage design, complete "
        "installed cost, lifecycle emissions, and new-discharge clearance "
        f"remained unresolved. Consequently, none of the {len(context['pareto'])} "
        "evaluated practical modes was evidence-supported; all remained conditional. "
        "The principal result is therefore a reproducible constraint-attrition "
        "framework that separates global burden-space potential from "
        "physically delivered heat and identifies exactly where practical "
        "claims fail."
    )
    _add_cited_paragraph(document, abstract)
    _add_cited_paragraph(
        document,
        (
            "Keywords: waste-heat recovery; thermal storage; heat stress; "
            "constraint waterfall; finite dispatch; evidence audit"
        ),
        bold_prefix="Keywords:",
    )

    document.add_heading("1. Introduction", level=1)
    _add_cited_paragraph(
        document,
        (
            "Anthropogenic and industrial heat is small in the global mean "
            "energy balance but can be spatially concentrated enough to alter "
            "regional atmospheric conditions.{1,2} This mismatch between a "
            "small global mean and large local fluxes motivates a question "
            "that is usually left implicit: if one unit of thermal energy "
            "must remain in the Earth system, where and when does it impose "
            "the least human thermal-stress burden?"
        ),
    )
    _add_cited_paragraph(
        document,
        (
            "Energy-systems research already treats heat as a spatial and "
            "temporal resource. Transportation formulations allocate thermal "
            "resources within district systems,{3} thermal storage supplies "
            "short- and long-duration flexibility,{4} and low-temperature "
            "networks integrate distributed sources and stores.{5} These "
            "applications operate at engineered network scales. A global "
            "analysis therefore cannot be interpreted as evidence that "
            "low-grade heat is practically transportable across continents. "
            "Its defensible role is to establish a conditional screening "
            "bound within a prescribed atmospheric slab and identify which "
            "constraints cause that bound to collapse."
        ),
    )
    _add_cited_paragraph(
        document,
        (
            "Optimal transport supplies a mathematical language for conserved "
            "allocation,{6} and Earth-system studies have used related "
            "Wasserstein methods for oceanographic fields, sea ice, and "
            "climate-model evaluation.{7-9} However, a diagnostic distance "
            "between fields is not a physical heat-transfer design, and the "
            "calculation implemented here is a one-unit marginal pair search "
            "rather than a finite-mass optimal-transport plan. The source, "
            "sink, travel distance, storage lag, capacity, and environmental "
            "exclusions must remain explicit."
        ),
    )
    _add_cited_paragraph(
        document,
        (
            "Human heat-stress indices differ materially in their physical "
            "content and empirical domains.{10} Global reanalysis products "
            "support consistent gridded thermal indices,{11} while observed "
            "humid-heat extremes and population exposure establish the "
            "importance of compound temperature and humidity.{12,13} We "
            "therefore use a transparent convex thermal-stress burden rather "
            "than a mortality endpoint and test alternative metrics and "
            "parameters."
        ),
    )
    _add_cited_paragraph(
        document,
        (
            "A persistent gap is the link between such theoretical allocation "
            "value and a real low-grade thermal-energy pathway. Source heat, "
            "temperature grade, coincident demand, route hydraulics, losses, "
            "auxiliary electricity, storage, cost, and environmental "
            "admissibility are often incomplete or reported at incompatible "
            "resolutions. Treating missing inputs as favorable assumptions "
            "would make the apparent practical value irreproducible."
        ),
    )
    _add_cited_paragraph(
        document,
        (
            "This study therefore makes a two-scale contribution. First, it "
            "compares same-time spatial relocation, same-location temporal "
            "storage, and joint spatiotemporal allocation within one "
            "conserved-energy screening objective. Second, it audits an "
            "empirical waste-heat source and a documented receiving pathway "
            "without inventing missing engineering inputs. Third, it uses a "
            "constraint waterfall and mode-decision analysis to identify how "
            "much theoretical value can be translated into a supported "
            "physical-energy claim. The central question is not whether a "
            "mathematical sink exists, but where the global-to-real evidence "
            "chain survives or fails (Fig. 1)."
        ),
    )
    if inline_assets:
        _add_inline_figure(
            document,
            "figure_1_framework.png",
            legends[0],
        )

    document.add_heading("2. Methods", level=1)
    document.add_heading("2.1 Data and harmonization", level=2)
    _add_cited_paragraph(
        document,
        (
            "The meteorological baseline used the NCEP/NCAR Reanalysis 1 "
            "architecture described by Kalnay et al.{14} We acquired complete "
            "2023 annual files for two-metre air temperature and specific "
            "humidity, surface pressure, 10-m winds, surface radiative fluxes, "
            "sea-ice fraction, snow water equivalent, and the land-sea mask. "
            f"The common Gaussian grid contained {context['latitude_count']} "
            f"latitudes, {context['longitude_count']} longitudes, and "
            f"{context['time_count']} six-hourly observations."
        ),
    )
    _add_cited_paragraph(
        document,
        (
            "Population was taken from the WorldPop 2020 unconstrained global "
            "grid, a high-resolution gridded population product consistent "
            "with published WorldPop methods.{15} It was conservatively "
            "aggregated to the climate grid; the resulting total was "
            f"{context['population'] / 1e9:.3f} billion. Natural Earth 1:110m "
            "Admin-0 boundaries supplied country labels. Every raw snapshot "
            "was stored locally with its source URL, retrieval time, file "
            "size, SHA-256 checksum, license terms, and completeness status."
            " The public inputs are summarized in Table 1."
        ),
    )
    if inline_assets:
        _add_table(document, *main_tables[0])

    document.add_heading("2.2 Thermal-stress burden", level=2)
    _add_cited_paragraph(
        document,
        (
            "Specific humidity q and pressure p define vapour pressure as "
            "e = qp/(0.622 + 0.378q). Humidex is H = T_C + "
            "(5/9)(e_hPa - 10). For reference level H0 and curvature γ, the "
            "stress excess is s = max(H - H0, 0), and cell-time burden is "
            "B = P s^γ, where P is population. The primary configuration used "
            "H0 = 26 and γ = 2. Air temperature and a validity-masked Stull "
            "wet-bulb proxy were sensitivity metrics; wet-bulb estimates were "
            "not extrapolated outside their documented temperature and "
            "relative-humidity domain.{16} Wet-bulb marginal burdens used a "
            "centered numerical temperature derivative with specific "
            "humidity and pressure held fixed."
        ),
    )
    _add_cited_paragraph(
        document,
        (
            "For grid-cell area A, mixing height h, air density ρ, and specific "
            "heat c_p, atmospheric heat capacity is C = Ahρc_p. A heat "
            "increment Q changes temperature by ΔT = Q/C. At γ = 2, the "
            "marginal burden per GJ is m = Pγs^(γ-1)10^9/C. This derivative is "
            "a burden index, not a monetary value, mortality estimate, or "
            "measure of total planetary heat."
        ),
    )

    document.add_heading("2.3 Conserved allocation problem", level=2)
    _add_cited_paragraph(
        document,
        (
            "A decision removes Q from source cell-time (i,t) and releases the "
            "same Q at sink cell-time (j,t′). Transport and storage appear as "
            "objective penalties rather than fictitious energy losses. For a "
            "one-GJ perturbation, the evaluated value is V = m_source - "
            "m_sink - c_d d - c_s(t′ - t), where d is great-circle distance, "
            "c_d is the transport penalty, and c_s is the storage penalty per "
            "six-hour step. No-relocation has value zero; negative feasible "
            "optima are retained rather than clipped. These coefficients are "
            "abstract conditional penalties, not calibrated costs, losses, "
            "emissions, or technology efficiencies."
        ),
    )
    _add_cited_paragraph(
        document,
        (
            "Spatial allocation used same-time source and sink states. "
            "Temporal allocation used the same cell at a later time. Joint "
            "allocation allowed both changes. A spherical cKDTree enumerated "
            "only cells within the maximum distance and evaluated exact "
            "great-circle distances for candidates, avoiding an infeasible "
            "global dense pairwise matrix. The canonical primary sink mask "
            "required land fraction ≥0.5 and excluded cryosphere cells, "
            "defined by sea-ice fraction ≥0.15 or snow water equivalent "
            "≥1 kg m^-2. Unrestricted masks were retained as diagnostics. "
            "Canonical assumptions are summarized in Table 2."
        ),
    )
    if inline_assets:
        _add_table(document, *main_tables[1])

    document.add_heading("2.4 Robustness, finite energy, and diagnostics", level=2)
    _add_cited_paragraph(
        document,
        (
            "The Humidex sensitivity crossed three reference levels, three "
            "curvatures, and three prescribed mixing heights. Air temperature "
            "and wet-bulb sensitivities each varied three reference levels at "
            "fixed curvature and mixing height. Progressive mask ablation "
            "added distance, cryosphere, and ocean exclusions. "
            "Finite-energy tests recomputed exact quadratic burden changes "
            "for 1, 10^3, 10^5, and 10^6 GJ and enforced both an energy cap "
            "and a maximum local temperature perturbation. Seasonal marginal "
            "quantiles used all eligible populated land source cell-times in "
            "each local day/night group; total burden summed all grid cells "
            "and times in that group."
        ),
    )
    _add_cited_paragraph(
        document,
        (
            "A reduced-order seven-day trajectory advected the selected "
            "release point using six-hourly 10-m winds and recorded a "
            "diffusion-radius diagnostic. It did not calculate plume "
            "concentration or thermal response and was used only to flag "
            "downstream classification changes. Surface "
            "longwave fields were retained as descriptive context. Observed "
            "radiation products such as CERES EBAF characterize the radiation "
            "budget but do not identify the incremental radiative response to "
            "added surface heat,{17} whereas engineered radiative cooling "
            "depends on device-specific spectral properties.{18} We therefore "
            "made no causal radiative-disposal claim."
        ),
    )

    document.add_heading("2.5 Auditable real thermal-energy case", level=2)
    _add_cited_paragraph(
        document,
        (
            "The practical audit used the public 2023 Frontier "
            "high-performance-computing facility workbook as the measured "
            "source. Ten-minute waste heat, coolant supply and return "
            "temperature, flow, compute power, accessory power, and total "
            "power were retained without filling missing timestamps. The "
            "receiving pathway used the ORNL waste-heat-recovery report and "
            "official metadata, campus-map, and environmental records. Figure "
            "13 of the report supplied measured annual average and maximum "
            "space-heating demand summaries for the selected buildings, but "
            "not a synchronized hourly demand trace."
        ),
    )
    _add_cited_paragraph(
        document,
        (
            "Each required case input was classified as FOUND_MEASURED, "
            "FOUND_DERIVABLE, FOUND_BOUND_ONLY, or NOT_FOUND. Practical "
            "pathways were separately classified as SUPPORTED, CONDITIONAL, "
            "or EXCLUDED. Measured facility accessory electricity was not "
            "substituted for incremental heat-recovery auxiliaries, existing "
            "steam-line distance and loss were not assigned to a new hot-water "
            "route, and generic storage or cost parameters were not presented "
            "as case design values (Table 4)."
        ),
    )
    if inline_assets:
        _add_table(document, *main_tables[3])

    document.add_heading(
        "2.6 Constraint waterfall, finite dispatch, and mode decision",
        level=2,
    )
    _add_cited_paragraph(
        document,
        (
            "The global-to-real waterfall retained burden-space and "
            "physical-energy quantities as separate domains. It evaluated the "
            "global theoretical bound, environmental admissibility, geographic "
            "locality, pathway eligibility, technology evidence, storage and "
            "capacity, demand evidence, finite dispatch, and the supported "
            "practical endpoint. No conversion factor was used between the "
            "thermal-stress burden index and delivered MWh."
        ),
    )
    _add_cited_paragraph(
        document,
        (
            "Finite dispatch used the measured 10-minute source series, "
            "explicit unavailable intervals, a report-derived heat-pump COP "
            "and capacity, constant lower and upper demand bounds, and no, "
            "formula-sized, or generic-reference storage. The dispatch "
            "preserved thermal balance at every interval. The mode analysis "
            "compared useful heat delivered, demand served, residual source "
            "rejection, heat-pump electricity, storage capacity, and storage "
            "losses. Cost and lifecycle CO2e were excluded from practical mode "
            "selection because case-specific evidence was incomplete."
        ),
    )

    document.add_heading("3. Results", level=1)
    document.add_heading("3.1 Global marginal-burden field", level=2)
    jja = context["seasonal"][
        (context["seasonal"]["season"] == "JJA")
        & (context["seasonal"]["local_period"] == "day")
    ].iloc[0]
    _add_cited_paragraph(
        document,
        (
            "The annual peak field was concentrated in densely populated "
            "humid regions (Supplementary Fig. S4). Upper-tail marginal burden varied "
            "seasonally but remained non-zero by both local day and night "
            f"(Supplementary Table S1 and Fig. S5). The JJA daytime 95th "
            "percentile was "
            f"{jja['marginal_burden_p95']:.1f} burden units per GJ."
        ),
    )
    _add_cited_paragraph(
        document,
        (
            f"The maximum canonical source occurred in "
            f"{spatial['source_country']} at "
            f"{spatial['source_latitude']:.2f}°N, "
            f"{spatial['source_longitude']:.2f}°E, with marginal benefit "
            f"{spatial['source_marginal_benefit']:.1f} per GJ. "
            f"{spatial['source_country']} was selected in all "
            f"{len(robustness)} scenarios in the canonical parameter sweep "
            "(Supplementary Table S2 and Fig. S1). "
            f"Across {len(structural)} prespecified structural scenarios, source "
            f"country labels included "
            f"{', '.join(sorted(structural['source_country'].unique()))}; "
            "theoretical value persisted while selected source and sink "
            "coordinates shifted (Supplementary Table S6)."
        ),
    )

    document.add_heading("3.2 Spatial, temporal, and joint bounds", level=2)
    _add_cited_paragraph(
        document,
        (
            f"With the land requirement and cryosphere exclusion but no "
            "transport penalty, "
            f"the spatial solution moved one GJ from "
            f"{spatial['source_country']} to a representative co-optimal "
            f"grid cell labeled {spatial['sink_country']} over "
            f"{spatial['distance_km']:.0f} km and reduced marginal burden by "
            f"{spatial['net_benefit']:.1f} units per GJ. At a transport "
            f"penalty of 10^-5 per GJ-km, the sink moved to "
            f"{small_penalty['sink_country']} at "
            f"{small_penalty['distance_km']:.0f} km with negligible loss of "
            "objective value (Fig. 2)."
        ),
    )
    if inline_assets:
        _add_inline_figure(
            document,
            "figure_2_primary_spatial_transfer.png",
            legends[1],
        )
    _add_cited_paragraph(
        document,
        (
            f"For the same peak source cell-time used by the spatial result, "
            f"a {context['temporal_lag_hours']}-hour delay at the source "
            "location produced "
            f"{context['matched_temporal']['net_benefit']:.1f} units per GJ, "
            f"or {context['temporal_fraction']:.1%} of the simultaneous "
            "spatial bound (Table 3; Fig. 3). With zero penalties, "
            "matched-event joint allocation selected the immediate spatial sink."
        ),
    )
    if inline_assets:
        _add_table(document, *main_tables[2])
        _add_inline_figure(
            document,
            "figure_3_allocation_mode_comparison.png",
            legends[2],
        )
    document.add_heading(
        "3.3 Exclusion masks, penalty sensitivity, and finite-energy behavior",
        level=2,
    )
    safety = context["safety"]
    _add_cited_paragraph(
        document,
        (
            f"{int(safety['sink_is_ocean'].sum())} of the {len(safety)} "
            "progressive mask scenarios selected representative ocean sinks. "
            "Cryosphere exclusion alone did not "
            "eliminate the selected ocean cell because it was not classified "
            "as cryosphere. Ocean exclusion changed the representative "
            "co-optimal sink geography to land (Supplementary Table S3 and "
            "Fig. S6). The identical "
            "zero-burden objective values reveal substantial sink "
            "non-uniqueness and show why zero population cannot serve as an "
            "environmental criterion."
        ),
    )
    costly_spatial = context["costly_spatial"]
    costly_joint = context["costly_joint"]
    _add_cited_paragraph(
        document,
        (
            f"At a transport penalty of 5 per GJ-km, the best spatial-only "
            f"objective was {costly_spatial['net_benefit']:.1f}, so "
            "no-relocation dominated. The joint model instead selected "
            f"same-location temporal transfer and retained "
            f"{costly_joint['net_benefit']:.1f} units per GJ. Across the full "
            "penalty grid, the selected mathematical candidate shifted from "
            "spatial transfer toward temporal transfer (Supplementary Fig. "
            "S7). Because the coefficients are uncalibrated, this is a "
            "conditional mode map rather than evidence for an engineering "
            "tradeoff."
        ),
    )
    feasible = finite_q[finite_q["status"] == "feasible"]
    _add_cited_paragraph(
        document,
        (
            f"Transfers through {feasible['energy_gj'].max():g} GJ were "
            "feasible and showed a small decline in benefit per GJ because "
            "finite cooling reduced the average source benefit; sink burden "
            "remained zero in the reported feasible cases. The "
            f"{infeasible['energy_gj'].iloc[0]:,.0f}-GJ case was infeasible "
            "under the 0.1-K local temperature cap (Supplementary Table S4 "
            "and Fig. S2). "
            "Thus, the one-unit result is a marginal bound and cannot be "
            "scaled indefinitely."
        ),
    )
    cryosphere_hours = context["redistribution"].loc[
        context["redistribution"]["nearest_cell_cryosphere"], "horizon_hours"
    ]
    _add_cited_paragraph(
        document,
        (
            "The reduced-order wind trajectory entered cryosphere-classified "
            "nearest cells after "
            f"{', '.join(f'{value:g}' for value in cryosphere_hours)} h "
            "(Supplementary Table S5 and Fig. S3). This point-trajectory "
            "encounter is only "
            "a flag for follow-up; it is not a heat-plume simulation or "
            "evidence of downstream impact."
        ),
    )

    document.add_heading("3.4 Measured source and receiving-demand evidence", level=2)
    values = context["manuscript_values"]
    _add_cited_paragraph(
        document,
        (
            "The Frontier source record contained "
            f"{values.loc['frontier_valid_intervals', 'display_value']} valid "
            "10-minute intervals and "
            f"{values.loc['frontier_missing_intervals', 'display_value']} "
            "unavailable intervals. Observed valid-interval waste heat totaled "
            f"{values.loc['frontier_observed_heat', 'display_value']} MWh-th. "
            "The measured distribution had median "
            f"{values.loc['frontier_waste_heat_median', 'display_value']} MW, "
            "95th percentile "
            f"{values.loc['frontier_waste_heat_p95', 'display_value']} MW, and "
            "maximum "
            f"{values.loc['frontier_waste_heat_max', 'display_value']} MW. "
            "Median supply and return temperatures were "
            f"{values.loc['frontier_supply_temperature_median', 'display_value']} "
            "°C and "
            f"{values.loc['frontier_return_temperature_median', 'display_value']} "
            "°C, respectively (Fig. 4)."
        ),
    )
    _add_cited_paragraph(
        document,
        (
            "The ORNL report documented average and maximum space-heating "
            "demand of "
            f"{values.loc['ornl_case_a_average', 'display_value']} and "
            f"{values.loc['ornl_case_a_maximum', 'display_value']} MW-th for "
            "the 5600-5700-5800 complex. The selected larger building set had "
            f"{values.loc['ornl_case_b_selected_average', 'display_value']} "
            "MW-th average and "
            f"{values.loc['ornl_case_b_selected_maximum', 'display_value']} "
            "MW-th maximum demand. These summaries bounded demand magnitude "
            "but did not reconstruct coincidence with the source series. Of "
            f"{len(context['case_evidence'])} required evidence items, "
            f"{int(context['case_evidence']['status'].eq('NOT_FOUND').sum())} "
            "were not found and "
            f"{int(context['case_evidence']['status'].eq('FOUND_BOUND_ONLY').sum())} "
            "were available only as bounds."
        ),
    )
    if inline_assets:
        _add_inline_figure(
            document,
            "figure_4_case_source_and_demand.png",
            legends[3],
        )

    document.add_heading("3.5 Global-to-real constraint attrition", level=2)
    _add_cited_paragraph(
        document,
        (
            "The burden-space branch retained the global marginal screening "
            "values but could not assign a practical environmental or "
            "technology conversion to them. The physical-energy branch began "
            "with the measured source and documented demand bounds, then "
            "failed to close route, capture, auxiliary, storage-design, "
            "complete-cost, lifecycle-emissions, and external-discharge gates "
            "(Fig. 5 and Table 5). The terminal supported endpoint was therefore zero "
            "practical modes rather than an attenuated numerical fraction of "
            "the global burden-space optimum."
        ),
    )
    if inline_assets:
        _add_inline_figure(
            document,
            "figure_5_constraint_waterfall.png",
            legends[4],
        )
    if inline_assets:
        _add_table(document, *main_tables[4])

    document.add_heading("3.6 Conditional finite dispatch and mode decision", level=2)
    _add_cited_paragraph(
        document,
        (
            "Across the nonzero-demand dispatch sensitivities, served fraction "
            "ranged from "
            f"{values.loc['dispatch_served_fraction_min', 'display_value']} to "
            f"{values.loc['dispatch_served_fraction_max', 'display_value']}. "
            "Residual source rejection ranged from "
            f"{values.loc['dispatch_residual_rejection_min', 'display_value']} "
            "to "
            f"{values.loc['dispatch_residual_rejection_max', 'display_value']} "
            "MWh-th. The maximum absolute timestep thermal-balance error was "
            f"{values.loc['dispatch_max_balance_error', 'display_value']} "
            "MWh-th. These results demonstrate numerical dispatch closure but "
            "not calibrated annual operation."
        ),
    )
    _add_cited_paragraph(
        document,
        (
            "All six compared reuse and storage modes were Pareto efficient "
            "within at least one demand basis because useful heat, residual "
            "rejection, electricity, storage capacity, and losses traded off. "
            "Direct local reuse was the least-assumption mode within each "
            "demand bound; storage increased demand served but introduced an "
            "unsupported case-design requirement. Every mode remained "
            "CONDITIONAL, and none was selected as a supported practical "
            "endpoint (Fig. 6 and Table 6)."
        ),
    )
    if inline_assets:
        _add_inline_figure(
            document,
            "figure_6_practical_modes.png",
            legends[5],
        )
    if inline_assets:
        _add_table(document, *main_tables[5])

    document.add_heading("4. Discussion", level=1)
    document.add_heading("4.1 Energy-system interpretation", level=2)
    _add_cited_paragraph(
        document,
        (
            "The global result establishes the opportunity created by "
            "spatiotemporal heterogeneity. For the peak source cell-time, an "
            f"{context['temporal_lag_hours']}-hour delay yielded "
            f"{context['temporal_fraction']:.0%} of the "
            "simultaneous spatial bound. This comparison makes temporal "
            "management commensurable with spatial allocation inside the "
            "screening model, but it remains a burden-space result rather than "
            "a heat-delivery efficiency."
        ),
    )
    _add_cited_paragraph(
        document,
        (
            "The real-case audit shows why a large theoretical opportunity "
            "does not automatically become a supported system result. Source "
            "heat and temperature grade were measured, demand magnitude and "
            "heat-pump performance were documented, and repeated dispatch was "
            "numerically feasible under explicit bounds. Yet the absence of a "
            "coincident demand trace and integrated route, auxiliary, storage, "
            "cost, lifecycle, and environmental design prevented selection of "
            "a supported mode. This negative constraint result is itself the "
            "applied finding: evidence attrition, not optimizer performance, "
            "determined practical claim eligibility."
        ),
    )
    document.add_heading("4.2 Environmental interpretation", level=2)
    _add_cited_paragraph(
        document,
        (
            "The unconstrained optimizer repeatedly preferred unpopulated "
            "ocean or polar cells because the objective measured human burden. "
            "This is a diagnostic failure mode, not a recommendation. Ocean, "
            "cryosphere, ecosystem, and Earth-system consequences require "
            "independent constraints. The trajectory experiment further shows "
            "that a stationary release-cell mask cannot guarantee that "
            "downstream heat remains outside sensitive environments."
        ),
    )
    _add_cited_paragraph(
        document,
        (
            "Likewise, the study reallocates heat rather than destroying it. "
            "No result should be interpreted as planetary cooling. A separate "
            "radiative-loss analysis would require a causal response model for "
            "added heat, not an observed longwave climatology. Device-scale "
            "radiative cooling literature does not by itself justify a "
            "planetary surface-flux derivative."
        ),
    )
    document.add_heading("4.3 Limitations and next steps", level=2)
    _add_cited_paragraph(
        document,
        (
            "The atmospheric control volume uses a prescribed mixing height "
            "and instantaneous linear temperature response. It does not "
            "resolve boundary-layer dynamics, land or ocean heat uptake, "
            "cloud feedbacks, precipitation, or circulation responses. "
            "Humidex is only one stress metric; the robustness and structural "
            "replication analyses reduce but do not remove metric dependence. "
            "The replication matrix spans multiple years, two related NCEP "
            "reanalysis products, two population surfaces, factor-two grid "
            "aggregation, and three thermal metrics. NCEP-DOE Reanalysis 2 "
            "shares lineage and resolution with NCEP/NCAR Reanalysis 1, so this "
            "is not an independent modern high-resolution replication. "
            "Population is static, and the burden function is not a "
            "health-impact function."
        ),
    )
    _add_cited_paragraph(
        document,
        (
            "The case study is limited by public evidence rather than by an "
            "attempt to estimate favorable missing values. A practical "
            "follow-up requires synchronized receiving demand, a surveyed "
            "route and hydraulic design, measured or guaranteed capture and "
            "auxiliary performance, storage sizing where applicable, complete "
            "installed cost with price year, lifecycle inventory, equipment "
            "availability, and route-specific environmental review. Until "
            "those inputs exist, dispatch, economics, emissions, and mode "
            "selection remain sensitivity analyses."
        ),
    )
    _add_cited_paragraph(
        document,
        (
            "The closest literatures already cover district-energy resource "
            "allocation, thermal storage, low-temperature networks, "
            "anthropogenic heat, heat-stress mapping, and conserved-allocation "
            "methods. The novelty claimed here is narrower: a reproducible "
            "two-scale chain that preserves conserved-energy accounting, "
            "diagnoses pathological global optima, and refuses to translate a "
            "theoretical bound into practical performance when required "
            "evidence is absent."
        ),
    )

    document.add_heading("5. Conclusions", level=1)
    _add_cited_paragraph(
        document,
        (
            "A heterogeneous thermal-stress field creates a large marginal "
            "screening value for moving heat, and short temporal delay "
            "preserves part of the matched spatial bound. The empirical "
            "Frontier-ORNL pathway confirms substantial measured low-grade "
            "source heat and documented receiving-demand bounds, but the "
            "constraint waterfall leaves no supported practical mode because "
            "critical coincident, engineering, economic, lifecycle, and "
            "environmental evidence is missing. Global burden-space value and "
            "physical-energy delivery must therefore remain separate. The "
            "main contribution is a reproducible method for showing how much "
            "of a theoretical heat-allocation opportunity survives real "
            "evidence gates—and for reporting zero when the gates do not close."
        ),
    )

    document.add_heading("Data and code availability", level=1)
    _add_cited_paragraph(
        document,
        (
            "All analysis code, configurations, processed result tables, "
            "checksums, and figure-generation scripts are available at "
            "https://github.com/bougtoir/global-heat-allocation. Large public "
            "raw files are not bundled with the manuscript archive; immutable "
            "source URLs, retrieval conditions, sizes, SHA-256 checksums, and "
            "acquisition scripts are provided so that they can be independently "
            "retrieved and verified. The repository must be archived and its "
            "permanent identifier inserted before submission."
        ),
    )
    document.add_heading("Funding", level=1)
    _add_cited_paragraph(
        document,
        "[To be completed and verified by the author before submission.]",
    )
    document.add_heading("Declaration of competing interest", level=1)
    _add_cited_paragraph(
        document,
        "[To be completed and verified by the author before submission.]",
    )
    document.add_heading(
        "Declaration of generative AI and AI-assisted technologies", level=1
    )
    _add_cited_paragraph(
        document,
        (
            "During preparation of this work, the author used Devin "
            "(Cognition AI) to assist with code generation, analysis workflow "
            "development, and language drafting. The author reviewed and "
            "edited all outputs, verified the underlying data and citations, "
            "and takes full responsibility for the content."
        ),
    )

    document.add_heading("References", level=1)
    for number, key in enumerate(REFERENCE_KEYS, start=1):
        document.add_paragraph(_format_reference(number, references[key]))

    if not inline_assets:
        document.add_section(WD_SECTION.NEW_PAGE)
        document.add_heading("Tables", level=1)
        for frame, caption in main_tables:
            _add_table(document, frame, caption)

        document.add_heading("Figure legends", level=1)
        for legend in legends:
            document.add_paragraph(legend)

    SUBMISSION.mkdir(parents=True, exist_ok=True)
    filename = (
        "manuscript_applied_energy_inline.docx"
        if inline_assets
        else "manuscript_applied_energy.docx"
    )
    output = SUBMISSION / filename
    document.save(output)
    return output


def _build_supplement(context: dict) -> Path:
    document = Document()
    _configure_document(document)
    document.add_heading(
        "Supplementary material: From global heat-allocation bounds to a real "
        "waste-heat pathway",
        level=0,
    )
    _add_cited_paragraph(
        document,
        (
            "This supplement is generated from the canonical result files. "
            "It reports complete diagnostic summaries and does not introduce "
            "additional empirical data."
        ),
    )
    document.add_heading("Supplementary methods", level=1)
    _add_cited_paragraph(
        document,
        (
            "The exact spatial search processes each six-hourly field in time "
            "order. Candidate sinks within the maximum great-circle distance "
            "are identified with a spherical cKDTree and rescored with exact "
            "haversine distance. Temporal candidates are vectorized over lag "
            "steps. Joint allocation compares spatial, temporal, and combined "
            "candidates under the same objective and eligibility masks."
        ),
    )
    _add_cited_paragraph(
        document,
        (
            "For finite Q and quadratic burden, source benefit is evaluated "
            "analytically as P(2sΔ - Δ²) when s > Δ and Ps² when 0 < s ≤ Δ. "
            "Sink burden is P(2sΔ + Δ²) when s ≥ 0 and "
            "P max(s + Δ, 0)² otherwise. These stable expressions avoid "
            "subtracting nearly equal large burden values."
        ),
    )
    _add_cited_paragraph(
        document,
        (
            "Validation checks scenario counts, objective identities, maximum "
            "distance, temporal same-location behavior, source and sink "
            "lookups against the processed cube, energy conservation, "
            "no-relocation retention, finite-Q monotonicity, seasonal group "
            "coverage, trajectory conservation, and explicit noncausal "
            "radiation interpretation."
        ),
    )
    document.add_heading("Supplementary tables", level=1)
    _add_table(
        document,
        context["seasonal"],
        "Table S1. Seasonal and local day-night marginal-burden summary.",
    )
    _add_table(
        document,
        pd.read_csv(TABLES / "table_6_robustness_summary.csv"),
        "Table S2. Source-hotspot robustness summary.",
    )
    _add_table(
        document,
        context["safety"],
        "Table S3. Progressive exclusion-mask ablation results.",
    )
    _add_table(
        document,
        context["finite_q"],
        "Table S4. Finite-transfer and local-capacity sensitivity.",
    )
    _add_table(
        document,
        context["redistribution"],
        "Table S5. Reduced-order natural-redistribution diagnostic.",
    )
    structural_columns = [
        "scenario_id",
        "climate_product",
        "climate_year",
        "population_product",
        "spatial_coarsening_factor",
        "thermal_metric",
        "source_country",
        "spatial_zero_penalty_sink_country",
        "spatial_zero_penalty_net_benefit_per_gj",
        "temporal_to_spatial_value_ratio",
        "source_shift_km_from_canonical",
        "zero_penalty_sink_shift_km_from_canonical",
    ]
    _add_table(
        document,
        context["structural"][structural_columns],
        "Table S6. Prespecified structural-replication results.",
    )
    document.add_heading("Supplementary figure legends", level=1)
    legends = [
        (
            "Figure S1. Source-hotspot selections in the limited one-year "
            "parameter sweep. Humidex varies threshold, curvature, and "
            "prescribed mixing height; alternative metrics vary threshold."
        ),
        (
            "Figure S2. Net benefit per GJ across finite transferred-energy "
            "levels; the dashed line identifies infeasibility under the "
            "local temperature cap."
        ),
        (
            "Figure S3. Seven-day reduced-order 10-m-wind trajectory from the "
            "selected release cell, with 24-, 72-, and 168-hour positions."
        ),
        (
            "Figure S4. Annual peak population-weighted marginal "
            "thermal-stress burden per GJ on the common global grid."
        ),
        (
            "Figure S5. Seasonal day-night variation in the 95th percentile "
            "of marginal thermal-stress burden."
        ),
        (
            "Figure S6. Progressive mask ablation showing how distance, "
            "cryosphere, and ocean exclusions change the representative sink "
            "among non-unique candidates."
        ),
        (
            "Figure S7. Selected allocation mode and loss from the "
            "zero-penalty optimum across abstract transport and storage "
            "penalty coefficients."
        ),
    ]
    for legend in legends:
        document.add_paragraph(legend)
    output = SUBMISSION / "supplementary_material.docx"
    document.save(output)
    return output


def _build_highlights(context: dict) -> Path:
    values = context["manuscript_values"]
    lines = [
        "Highlights",
        "",
        "• Conserved heat allocation defines a global marginal screening bound",
        (
            f"• A peak-event {context['temporal_lag_hours']}-hour delay yields "
            f"{context['temporal_fraction']:.0%} of its spatial bound"
        ),
        (
            "• Measured Frontier source heat totals "
            f"{values.loc['frontier_observed_heat', 'display_value']} MWh-th"
        ),
        "• A constraint waterfall separates burden-space and delivered heat",
        "• No practical mode passes all case-specific evidence gates",
        "",
    ]
    output = SUBMISSION / "highlights.txt"
    output.write_text("\n".join(lines), encoding="utf-8")
    return output


def _build_cover_letter(context: dict) -> Path:
    document = Document()
    _configure_document(document)
    warning = document.add_paragraph()
    warning.add_run(
        "DRAFT—NOT FOR SUBMISSION UNTIL THE SCIENTIFIC AND ADMINISTRATIVE GATES "
        "IN THE CHECKLIST ARE RESOLVED"
    ).bold = True
    document.add_paragraph("Editor-in-Chief")
    document.add_paragraph("Applied Energy")
    document.add_paragraph()
    document.add_paragraph("Dear Editor,")
    _add_cited_paragraph(
        document,
        (
            "Please consider the manuscript “From global heat-allocation "
            "bounds to a real waste-heat pathway: a two-scale constraint "
            "analysis” as a Full Length Research Article in Applied Energy."
        ),
    )
    _add_cited_paragraph(
        document,
        (
            "The manuscript connects a conserved-heat global screening model "
            "to measured Frontier supercomputer waste heat and documented "
            "Oak Ridge National Laboratory heating-demand bounds. Its matched "
            "peak-event result is that a "
            f"{context['temporal_lag_hours']}-hour local delay "
            f"yields {context['temporal_fraction']:.0%} of the simultaneous "
            "spatial bound. A global-to-real constraint waterfall and finite "
            "dispatch then show that every evaluated practical mode remains "
            "conditional because essential case evidence is missing."
        ),
    )
    _add_cited_paragraph(
        document,
        (
            "The contribution is the transparent constraint-attrition result, "
            "not a claim that planetary heat transport or a calibrated "
            "Frontier-to-ORNL project is presently supported. Measured, "
            "derived, bounded, missing, conditional, and excluded quantities "
            "remain explicit; negative and infeasible outcomes are retained."
        ),
    )
    _add_cited_paragraph(
        document,
        (
            "This manuscript is original, is not under consideration "
            "elsewhere, and has been approved by all authors. "
            "[The author must verify this statement before submission.]"
        ),
    )
    _add_cited_paragraph(
        document,
        (
            "Suggested reviewers and any opposed reviewers will be supplied "
            "through the submission system after conflict-of-interest review."
        ),
    )
    document.add_paragraph("Sincerely,")
    document.add_paragraph("Tatsuki Onishi")
    document.add_paragraph("[Affiliation to be confirmed]")
    document.add_paragraph("bougtoir@gmail.com")
    output = SUBMISSION / "cover_letter.docx"
    document.save(output)
    return output


def _build_graphical_abstract(context: dict) -> tuple[Path, Path]:
    presentation = Presentation()
    presentation.slide_width = PptxInches(13.333)
    presentation.slide_height = PptxInches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    title = slide.shapes.add_textbox(
        PptxInches(0.45),
        PptxInches(0.25),
        PptxInches(12.4),
        PptxInches(0.65),
    )
    paragraph = title.text_frame.paragraphs[0]
    paragraph.text = "How much global heat-allocation value survives real constraints?"
    paragraph.alignment = PP_ALIGN.CENTER
    paragraph.runs[0].font.size = PptxPt(26)
    paragraph.runs[0].font.bold = True
    boxes = [
        (
            0.6,
            "Global screening bound",
            "conserved one-unit allocation\nburden space",
            "D9EAF2",
        ),
        (
            4.75,
            "Measured real case",
            "Frontier waste heat\nORNL demand bounds",
            "E8E1F2",
        ),
        (
            8.9,
            "Evidence waterfall",
            "dispatch | route | cost |\nlifecycle | environment",
            "DFF0E5",
        ),
    ]
    for left, heading, body, fill in boxes:
        shape = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
            PptxInches(left),
            PptxInches(1.35),
            PptxInches(3.75),
            PptxInches(2.0),
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = PptxRGBColor.from_string(fill)
        shape.line.color.rgb = PptxRGBColor(0, 114, 178)
        frame = shape.text_frame
        frame.clear()
        first = frame.paragraphs[0]
        first.text = heading
        first.alignment = PP_ALIGN.CENTER
        first.runs[0].font.size = PptxPt(19)
        first.runs[0].font.bold = True
        second = frame.add_paragraph()
        second.text = body
        second.alignment = PP_ALIGN.CENTER
        second.runs[0].font.size = PptxPt(14)
    for start, end in [(4.35, 4.75), (8.5, 8.9)]:
        connector = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            PptxInches(start),
            PptxInches(2.35),
            PptxInches(end),
            PptxInches(2.35),
        )
        connector.line.color.rgb = PptxRGBColor(80, 80, 80)
        connector.line.width = PptxPt(2)
    metrics = [
        (
            f"{context['spatial']['net_benefit']:.0f}",
            "matched spatial screening\nburden units per GJ",
            "0072B2",
        ),
        (
            context["manuscript_values"].loc["frontier_observed_heat", "display_value"],
            "observed Frontier source\nMWh-th",
            "009E73",
        ),
        (
            "0",
            "supported practical modes\nunder current evidence",
            "D55E00",
        ),
    ]
    for index, (value, label, color) in enumerate(metrics):
        left = 1.05 + index * 4.15
        circle = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.OVAL,
            PptxInches(left),
            PptxInches(4.05),
            PptxInches(1.45),
            PptxInches(1.45),
        )
        circle.fill.solid()
        circle.fill.fore_color.rgb = PptxRGBColor.from_string(color)
        circle.line.color.rgb = PptxRGBColor.from_string(color)
        circle.text = value
        circle.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        circle.text_frame.paragraphs[0].runs[0].font.size = PptxPt(21)
        circle.text_frame.paragraphs[0].runs[0].font.bold = True
        circle.text_frame.paragraphs[0].runs[0].font.color.rgb = PptxRGBColor(
            255, 255, 255
        )
        label_box = slide.shapes.add_textbox(
            PptxInches(left + 1.6),
            PptxInches(4.2),
            PptxInches(2.15),
            PptxInches(1.2),
        )
        label_paragraph = label_box.text_frame.paragraphs[0]
        label_paragraph.text = label
        label_paragraph.runs[0].font.size = PptxPt(14)
    footer = slide.shapes.add_textbox(
        PptxInches(0.7),
        PptxInches(6.55),
        PptxInches(12.0),
        PptxInches(0.5),
    )
    footer_paragraph = footer.text_frame.paragraphs[0]
    footer_paragraph.text = (
        "Burden-space potential and physical-energy delivery remain separate"
    )
    footer_paragraph.alignment = PP_ALIGN.CENTER
    footer_paragraph.runs[0].font.size = PptxPt(16)
    footer_paragraph.runs[0].font.italic = True
    pptx_output = SUBMISSION / "graphical_abstract_editable.pptx"
    presentation.save(pptx_output)

    figure, axis = plt.subplots(figsize=(13.33, 7.5))
    axis.set_axis_off()
    axis.set_title(
        "How much global heat-allocation value survives real constraints?",
        fontsize=22,
        weight="bold",
        pad=18,
    )
    for left, heading, body, fill in boxes:
        x_position = (left + 1.875) / 13.333
        axis.text(
            x_position,
            0.68,
            f"{heading}\n\n{body}",
            ha="center",
            va="center",
            transform=axis.transAxes,
            fontsize=13,
            bbox={
                "boxstyle": "round,pad=0.7",
                "facecolor": f"#{fill}",
                "edgecolor": "#0072b2",
            },
        )
    for x_position in [0.34, 0.65]:
        axis.annotate(
            "",
            xy=(x_position + 0.04, 0.68),
            xytext=(x_position - 0.02, 0.68),
            xycoords=axis.transAxes,
            arrowprops={"arrowstyle": "->", "lw": 2, "color": "#555555"},
        )
    for index, (value, label, color) in enumerate(metrics):
        axis.text(
            0.17 + index * 0.32,
            0.30,
            value,
            ha="center",
            va="center",
            transform=axis.transAxes,
            fontsize=19,
            color="white",
            weight="bold",
            bbox={
                "boxstyle": "circle,pad=0.7",
                "facecolor": f"#{color}",
                "edgecolor": f"#{color}",
            },
        )
        axis.text(
            0.17 + index * 0.32,
            0.13,
            label,
            ha="center",
            transform=axis.transAxes,
            fontsize=12,
        )
    axis.text(
        0.5,
        0.01,
        "Burden-space potential and physical-energy delivery remain separate",
        ha="center",
        transform=axis.transAxes,
        fontsize=13,
        style="italic",
    )
    png_output = SUBMISSION / "graphical_abstract.png"
    figure.savefig(png_output, dpi=300, bbox_inches="tight", facecolor="white")
    plt.close(figure)
    return pptx_output, png_output


def _build_editable_tables(context: dict) -> Path:
    document = Document()
    _configure_document(document)
    document.add_heading("Editable main tables", level=0)
    for index, (frame, caption) in enumerate(_main_tables(context)):
        _add_table(document, frame, caption)
        if index < 2:
            document.add_page_break()
    output = SUBMISSION / "tables_editable.docx"
    document.save(output)
    return output


def _phase_d_context() -> dict:
    energy = pd.read_csv(TECHNO_ECONOMIC / "energy_emissions.csv")
    economics = pd.read_csv(TECHNO_ECONOMIC / "economics.csv")
    sensitivity = pd.read_csv(TECHNO_ECONOMIC / "sensitivity.csv")
    frontier = pd.read_csv(TECHNO_ECONOMIC / "frontier_source_potential.csv").set_index(
        "metric"
    )
    parameters = pd.read_csv(TECHNO_ECONOMIC_PARAMETERS)

    case_energy = energy.loc[
        energy["scenario_id"].eq("C_case_cop_SRTV_total_output")
    ].iloc[0]
    generic_economics = economics.loc[
        economics["scenario_id"].eq("C_generic_full_system")
    ].iloc[0]
    missing_required = parameters.loc[
        parameters["value_kind"].eq("missing_required"),
        ["parameter_id", "parameter_name", "limitation"],
    ].copy()
    return {
        "case_energy": case_energy,
        "generic_economics": generic_economics,
        "sensitivity_count": len(sensitivity),
        "avoided_emissions_min": float(
            sensitivity["avoided_operational_co2e_kg_per_mwh"].min()
        ),
        "avoided_emissions_max": float(
            sensitivity["avoided_operational_co2e_kg_per_mwh"].max()
        ),
        "incremental_npc_min": float(sensitivity["incremental_npc_usd2025"].min()),
        "incremental_npc_max": float(sensitivity["incremental_npc_usd2025"].max()),
        "frontier_valid_observations": int(
            frontier.loc["valid_waste_heat_observations", "value"]
        ),
        "frontier_observed_energy_mwh": float(
            frontier.loc["observed_waste_heat_energy", "value"]
        ),
        "frontier_missing_timestamps": int(frontier.loc["missing_timestamps", "value"]),
        "missing_required": missing_required,
    }


def _phase_e_context() -> dict:
    parameters = pd.read_csv(ENVIRONMENTAL_PARAMETERS)
    receiving = pd.read_csv(
        ENVIRONMENTAL_CONSTRAINTS / "receiving_class_constraints.csv"
    )
    pathways = pd.read_csv(
        ENVIRONMENTAL_CONSTRAINTS / "frontier_pathway_assessment.csv"
    )
    storage = pd.read_csv(
        ENVIRONMENTAL_CONSTRAINTS / "storage_capacity_sensitivity.csv"
    )
    missing_required = parameters.loc[
        parameters["value_kind"].eq("missing_required"),
        ["parameter_id", "parameter_name", "limitation"],
    ].copy()
    external = pathways.loc[pathways["new_external_thermal_discharge"].eq("yes")]
    return {
        "receiving_class_count": len(receiving),
        "external_pathway_count": len(external),
        "excluded_external_pathway_count": int(
            external["assessment_status"].eq("excluded").sum()
        ),
        "storage_sensitivity_count": len(storage),
        "storage_capacity_min": float(storage["capacity_mwh_th"].min()),
        "storage_capacity_max": float(storage["capacity_mwh_th"].max()),
        "missing_required": missing_required,
    }


def _phase_f_context() -> dict:
    structural = pd.read_csv(STRUCTURAL_REPLICATION)
    return {
        "scenario_count": len(structural),
        "positive_spatial_count": int(
            structural["theoretical_spatial_value_positive"].sum()
        ),
        "nonunique_sink_count": int(
            structural["sink_nonunique_within_tolerance"].sum()
        ),
        "pathological_sink_count": int(
            structural["unrestricted_pathological_sinks_present"].sum()
        ),
        "positive_temporal_count": int(
            structural["temporal_net_benefit_per_gj"].gt(0.0).sum()
        ),
        "high_penalty_collapse_count": int(
            structural["spatial_value_collapses_at_high_penalty"].sum()
        ),
        "source_countries": sorted(structural["source_country"].unique()),
        "maximum_source_shift_km": float(
            structural["source_shift_km_from_canonical"].max()
        ),
        "maximum_sink_shift_km": float(
            structural["zero_penalty_sink_shift_km_from_canonical"].max()
        ),
    }


def _phase_g_context() -> dict:
    summary = pd.read_csv(DISPATCH / "finite_dispatch_summary.csv").set_index(
        "scenario_id"
    )
    if summary["practical_case_dispatch_claim_permitted"].any():
        raise ValueError("Sensitivity dispatch must not permit a practical case claim.")
    return {
        "scenario_count": len(summary),
        "interval_count": int(summary["interval_count"].iloc[0]),
        "zero_demand": summary.loc["zero_demand_rejection_baseline"],
        "direct_lower": summary.loc["direct_lower_demand_bound"],
        "direct_upper": summary.loc["direct_upper_demand_bound"],
        "minimum_lower": summary.loc["minimum_formula_storage_lower_demand"],
        "minimum_upper": summary.loc["minimum_formula_storage_upper_demand"],
        "generic_lower": summary.loc["generic_storage_lower_demand"],
        "generic_upper": summary.loc["generic_storage_upper_demand"],
        "maximum_thermal_balance_error_mwh": float(
            summary["maximum_absolute_thermal_balance_error_mwh"].max()
        ),
    }


def _build_nogo_report(context: dict) -> Path:
    document = Document()
    _configure_document(document)
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Submission Readiness Assessment: NO-GO")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(
        "From global heat-allocation bounds to a real waste-heat pathway: "
        "a two-scale constraint analysis"
    ).italic = True

    status = document.add_paragraph()
    status.alignment = WD_ALIGN_PARAGRAPH.CENTER
    status.paragraph_format.space_before = Pt(12)
    status_run = status.add_run("DECISION: NO-GO FOR SUBMISSION")
    status_run.bold = True
    status_run.font.size = Pt(16)
    status_run.font.color.rgb = RGBColor(192, 0, 0)

    allocations = context["allocations"]
    relocation_positive = int(allocations["relocation_dominates"].sum())
    scenario_count = len(allocations)
    matched_spatial = context["matched_spatial"]
    matched_temporal = context["matched_temporal"]
    phase_d = _phase_d_context()
    phase_e = _phase_e_context()
    phase_f = _phase_f_context()
    phase_g = _phase_g_context()
    case_energy = phase_d["case_energy"]
    generic_economics = phase_d["generic_economics"]

    document.add_heading("1. Executive assessment", level=1)
    document.add_paragraph(
        "The study formulates a global marginal-allocation problem in which heat "
        "is conserved and redistributed across space and time to reduce a "
        "population-weighted thermal-stress burden, then connects that theoretical "
        "screening bound to a measured Frontier waste-heat stream and an ORNL "
        "receiving pathway. Technology comparators and transparent techno-economic "
        "and operational-emissions accounting are now implemented. A conservative "
        "environmental constraint model also excludes unsupported external thermal "
        "sinks. A prespecified structural-replication matrix tests year, related "
        "reanalysis product, population surface, grid aggregation, and thermal "
        "metric choices. Seven finite repeated-dispatch sensitivities now conserve "
        "energy at each 10-minute timestep while exposing storage losses, unmet "
        "demand, and residual rejection. These additions materially strengthen the "
        "project, but they do not establish measured receiving demand, a complete "
        "project cost, lifecycle emissions, case-specific cumulative environmental "
        "load, or a case-specific integrated engineering design. Submission to "
        "Applied Energy should therefore not proceed at this stage."
    )
    document.add_paragraph(
        "This decision does not reject the research idea or mathematical framework. "
        "It reflects the fixed evidence rule: a gate is closed only by auditable "
        "data and reproducible analysis, not by a scenario assumption or narrative "
        "qualification."
    )

    document.add_heading("2. Scientific contribution established to date", level=1)
    contributions = [
        (
            "Formulates conserved heat as spatial relocation, temporal storage, "
            "and joint spatiotemporal allocation within one marginal objective."
        ),
        (
            "Provides an exact, memory-bounded global one-unit search without "
            "constructing an infeasible dense global pair matrix."
        ),
        (
            "Diagnoses how distance, land-ocean and cryosphere masks, local "
            "capacity, finite heat quantities, and abstract transport or storage "
            "penalties alter the mathematical optimum."
        ),
        (
            "Shows that low population does not imply a safe heat sink and "
            "separates mathematical optima from environmental or engineering "
            "recommendations."
        ),
        (
            "Integrates public inputs, acquisition ledgers, checksums, analysis "
            "code, and artifact generation into a reproducible screening pipeline."
        ),
        (
            "Selects a public measured waste-heat case and preserves missing "
            "timestamps rather than imputing unsupported annual energy."
        ),
        (
            "Separates source capture, heat-pump electricity, storage and delivery "
            "losses, auxiliaries, residual rejection, avoided fuel, cost, and "
            "operational emissions while preventing COP double counting."
        ),
        (
            "Replaces population-only sink safety with source-traced environmental "
            "receiving classes, regulatory criteria, physical storage sensitivity, "
            "and categorical exclusion of unsupported external discharge."
        ),
        (
            "Replicates the central structural findings across prespecified years, "
            "related reanalysis products, population surfaces, grid aggregation, "
            "and thermal metrics while reporting geographic instability."
        ),
        (
            "Implements auditable finite repeated dispatch with source availability, "
            "storage state, power limits, losses, auxiliaries, unmet demand, residual "
            "rejection, and explicit timestep energy-balance checks."
        ),
    ]
    for contribution in contributions:
        document.add_paragraph(contribution, style="List Number")

    document.add_heading("3. Current evidence", level=1)
    document.add_paragraph(
        f"The canonical analysis generated {scenario_count} scenarios; relocation "
        f"dominated no relocation in {relocation_positive}. The representative "
        f"matched-event spatial bound was "
        f"{matched_spatial['net_benefit']:.1f} burden units/GJ. A same-location "
        f"{context['temporal_lag_hours']}-hour delay yielded "
        f"{matched_temporal['net_benefit']:.1f} burden units/GJ, or "
        f"{context['temporal_fraction']:.1%} of that event-specific spatial bound."
    )
    document.add_paragraph(
        "The matched spatial result had "
        f"{int(matched_spatial['cooptimal_candidate_count'])} exactly co-optimal "
        "candidates and "
        f"{int(matched_spatial['near_optimal_candidate_count'])} candidates within "
        "1% of the optimum. The representative sink is therefore not unique and "
        "must not be interpreted as a policy recommendation."
    )
    document.add_paragraph(
        f"Phase F evaluated {phase_f['scenario_count']} prespecified structural "
        f"scenarios. Positive constrained spatial value, sink non-uniqueness, "
        f"pathological unrestricted sinks, positive temporal value, and "
        f"high-penalty spatial collapse each occurred in "
        f"{phase_f['positive_spatial_count']}/"
        f"{phase_f['scenario_count']}, "
        f"{phase_f['nonunique_sink_count']}/"
        f"{phase_f['scenario_count']}, "
        f"{phase_f['pathological_sink_count']}/"
        f"{phase_f['scenario_count']}, "
        f"{phase_f['positive_temporal_count']}/"
        f"{phase_f['scenario_count']}, and "
        f"{phase_f['high_penalty_collapse_count']}/"
        f"{phase_f['scenario_count']} scenarios, respectively. Source-country "
        f"labels included {', '.join(phase_f['source_countries'])}; maximum source "
        f"and zero-penalty sink coordinate shifts were "
        f"{phase_f['maximum_source_shift_km']:.0f} km and "
        f"{phase_f['maximum_sink_shift_km']:.0f} km."
    )
    document.add_paragraph(
        "The Frontier workbook provides "
        f"{phase_d['frontier_valid_observations']:,} valid waste-heat observations "
        f"totaling {phase_d['frontier_observed_energy_mwh']:,.1f} observed MWh-th. "
        f"{phase_d['frontier_missing_timestamps']:,} expected ten-minute timestamps "
        "are absent. No missing energy is imputed, and the observed source potential "
        "is not presented as delivered useful heat."
    )
    document.add_paragraph(
        f"Phase D evaluates {phase_d['sensitivity_count']:,} sensitivity cases. "
        "For the normalized case-C energy-only calculation using the eGRID SRTV "
        "total-output factor, avoided operational emissions are "
        f"{case_energy['avoided_operational_co2e_kg']:.2f} kg CO2e per MWh useful "
        "heat. Across the sensitivity grid, avoided operational emissions range "
        f"from {phase_d['avoided_emissions_min']:.2f} to "
        f"{phase_d['avoided_emissions_max']:.2f} kg CO2e per MWh, and incremental "
        f"NPC ranges from USD {phase_d['incremental_npc_min']:,.0f} to USD "
        f"{phase_d['incremental_npc_max']:,.0f}. Negative outcomes are retained."
    )
    document.add_paragraph(
        "The generic full-system comparator has an LCOH of USD "
        f"{generic_economics['lcoh_usd2025_per_mwh_th']:.2f}/MWh-th. This is not a "
        "Frontier project estimate. Case-specific LCOH, NPC, annual delivery, total "
        "residual rejection, and lifecycle CO2e remain intentionally withheld."
    )
    document.add_paragraph(
        f"Phase E defines {phase_e['receiving_class_count']} environmental receiving "
        "classes and excludes all "
        f"{phase_e['excluded_external_pathway_count']} modeled new external "
        "thermal-discharge alternatives. The generic source-backed tank sensitivity "
        f"contains {phase_e['storage_sensitivity_count']} cases spanning "
        f"{phase_e['storage_capacity_min']:.1f} to "
        f"{phase_e['storage_capacity_max']:.1f} MWh-th. This range is not a "
        "Frontier storage design; case volume and operating temperatures remain "
        "missing."
    )
    document.add_paragraph(
        f"Phase G evaluates {phase_g['scenario_count']} prespecified dispatch "
        f"sensitivities over {phase_g['interval_count']:,} ten-minute intervals. "
        "The zero-demand baseline rejects "
        f"{phase_g['zero_demand']['residual_source_rejection_mwh']:,.1f} MWh-th. "
        "The generic-storage lower- and upper-demand sensitivities serve "
        f"{phase_g['generic_lower']['demand_served_fraction']:.2%} and "
        f"{phase_g['generic_upper']['demand_served_fraction']:.2%} of constant "
        "reported demand bounds, respectively, while retaining "
        f"{phase_g['generic_lower']['residual_source_rejection_mwh']:,.1f} and "
        f"{phase_g['generic_upper']['residual_source_rejection_mwh']:,.1f} MWh-th "
        "of residual rejection. Maximum absolute timestep thermal-balance error is "
        f"{phase_g['maximum_thermal_balance_error_mwh']:.2e} MWh. These are "
        "sensitivities, not measured Frontier/ORNL operation."
    )

    document.add_heading("4. Progress since the prior NO-GO assessment", level=1)
    progress = pd.DataFrame(
        [
            {
                "Area": "Real heat case",
                "Completed evidence": (
                    "Measured Frontier source profile and an auditable ORNL "
                    "receiving pathway selected"
                ),
                "Evidence still withheld": (
                    "Synchronized hourly receiving demand and finite delivery"
                ),
            },
            {
                "Area": "Technology",
                "Completed evidence": (
                    "Case COP evidence and generic heat-pump, storage, and network "
                    "comparators"
                ),
                "Evidence still withheld": (
                    "Case capture loss, auxiliaries, route, hydraulics, and "
                    "balance-of-plant design"
                ),
            },
            {
                "Area": "Economics and emissions",
                "Completed evidence": (
                    "Energy balance, operational CO2e, CAPEX/OPEX, replacement, "
                    "discounting, and broad sensitivity"
                ),
                "Evidence still withheld": (
                    "Frontier installed cost, transport CAPEX, case LCOH/NPC, and "
                    "lifecycle CO2e"
                ),
            },
            {
                "Area": "Environmental receiving capacity",
                "Completed evidence": (
                    "Source-traced receiving classes, Tennessee thermal criteria, "
                    "external-sink exclusions, and generic tank sensitivity"
                ),
                "Evidence still withheld": (
                    "Case storage design, route clearance, finite residual load, "
                    "and any new outfall-specific permit and ecological evidence"
                ),
            },
            {
                "Area": "Finite dispatch",
                "Completed evidence": (
                    "Ten-minute source, demand-bound, storage, loss, auxiliary, "
                    "unmet-demand, and residual-rejection accounting"
                ),
                "Evidence still withheld": (
                    "Measured ORNL demand and case capture, storage, route, outage, "
                    "and environmental-capacity inputs"
                ),
            },
            {
                "Area": "Reproducibility",
                "Completed evidence": (
                    "Raw snapshots, checksums, parameter tables, result tables, "
                    "value ledger, regenerated manuscript package, validation, "
                    "rendering checks, and tests"
                ),
                "Evidence still withheld": (
                    "Permanent repository DOI and author-controlled submission metadata"
                ),
            },
            {
                "Area": "Decision evidence",
                "Completed evidence": (
                    "Global-to-real waterfall, six conditional practical modes, "
                    "post-freeze novelty audit, and adversarial manuscript review"
                ),
                "Evidence still withheld": (
                    "A supported practical endpoint under case-specific inputs"
                ),
            },
        ]
    )
    _add_table(document, progress, "Table 1. Evidence added after the initial audit.")

    document.add_heading("5. Mandatory gate status", level=1)
    gates = pd.DataFrame(
        [
            {
                "Gate": "1. Real source and receiving pathway",
                "Status": "PARTIAL—OPEN",
                "Blocking evidence": (
                    "Measured source exists; hourly receiving demand and synchronized "
                    "operation do not"
                ),
                "Required closure": (
                    "Auditable receiving profile supporting finite dispatch"
                ),
            },
            {
                "Gate": "2. Technology calibration",
                "Status": "PARTIAL—OPEN",
                "Blocking evidence": (
                    "Generic comparators cannot substitute for case capture, "
                    "auxiliary, route, and hydraulic measurements"
                ),
                "Required closure": (
                    "Case-specific design inputs or explicit pathway exclusion"
                ),
            },
            {
                "Gate": "3. Finite operation and dispatch",
                "Status": "PARTIAL—OPEN",
                "Blocking evidence": (
                    "Repeated sensitivity dispatch is complete, but demand is a "
                    "constant published bound rather than a measured load trace"
                ),
                "Required closure": (
                    "Measured receiving demand and case-specific design inputs"
                ),
            },
            {
                "Gate": "4. Losses, cost, and emissions",
                "Status": "PARTIAL—OPEN",
                "Blocking evidence": (
                    "Comparator analysis exists; case project cost, residual "
                    "rejection, and lifecycle boundary remain incomplete"
                ),
                "Required closure": (
                    "Case LCOH/NPC and supported operational or lifecycle claims"
                ),
            },
            {
                "Gate": "5. Environmental receiving capacity",
                "Status": "PARTIAL—OPEN",
                "Blocking evidence": (
                    "Sensitivity residual load is quantified, but case storage, "
                    "route, and permitted receiving capacity remain unresolved"
                ),
                "Required closure": (
                    "Case-specific cumulative load and receiving-design evidence"
                ),
            },
            {
                "Gate": "6. Structural uncertainty",
                "Status": "CLOSED WITH SCOPE LIMITATIONS",
                "Blocking evidence": (
                    "Nine prespecified scenarios cover year, related reanalysis "
                    "product, population surface, grid aggregation, and metric"
                ),
                "Required closure": (
                    "Retain scope statement: no modern independent "
                    "high-resolution reanalysis replication"
                ),
            },
            {
                "Gate": "7. Global-to-real constraint waterfall",
                "Status": "CLOSED—NEGATIVE RESULT",
                "Blocking evidence": (
                    "Burden-space and physical-energy branches have no supported "
                    "unit bridge"
                ),
                "Required closure": (
                    "Maintain separate units and the zero supported endpoint"
                ),
            },
            {
                "Gate": "8. Pareto and mode decision",
                "Status": "CLOSED WITH CONDITIONAL MODES",
                "Blocking evidence": (
                    "All six evaluated modes depend on incomplete case inputs"
                ),
                "Required closure": (
                    "Do not identify a supported or preferred practical mode"
                ),
            },
            {
                "Gate": "9. Novelty and journal fit",
                "Status": "CLOSED—NO-GO",
                "Blocking evidence": (
                    "The applied-energy implementation claim remains unsupported"
                ),
                "Required closure": (
                    "Acquire integrated case evidence or reconsider journal fit"
                ),
            },
            {
                "Gate": "10. Traceability and reproducibility",
                "Status": "CLOSED FOR CURRENT PACKAGE",
                "Blocking evidence": (
                    "A permanent repository DOI remains an administrative task"
                ),
                "Required closure": (
                    "Preserve the frozen ledger and regenerated package"
                ),
            },
            {
                "Gate": "11. Adversarial scientific review",
                "Status": "CLOSED—NO-GO RETAINED",
                "Blocking evidence": (
                    "No manuscript defect justifies inventing the missing case data"
                ),
                "Required closure": (
                    "Keep the open evidence gates explicit in all submission files"
                ),
            },
            {
                "Gate": "12. Health-claim control",
                "Status": "CLOSED—MAINTAIN",
                "Blocking evidence": (
                    "The endpoint must remain a thermal-stress burden proxy"
                ),
                "Required closure": "Continue excluding causal health claims",
            },
            {
                "Gate": "13. Submission administration",
                "Status": "OPEN",
                "Blocking evidence": (
                    "Final authorship, declarations, approvals, and repository DOI "
                    "are unconfirmed"
                ),
                "Required closure": "Author confirmation of all submission metadata",
            },
        ]
    )
    _add_table(document, gates, "Table 2. Evidence-based gate status.")

    document.add_heading("6. Case-specific evidence still missing", level=1)
    document.add_paragraph(
        "The Phase D and Phase E parameter registries preserve the following items "
        "as missing-required rather than substituting generic values:"
    )
    for row in phase_d["missing_required"].itertuples(index=False):
        document.add_paragraph(
            f"{row.parameter_name}: {row.limitation}",
            style="List Bullet",
        )
    for row in phase_e["missing_required"].itertuples(index=False):
        document.add_paragraph(
            f"{row.parameter_name}: {row.limitation}",
            style="List Bullet",
        )

    document.add_heading("7. Minimum conditions for a GO decision", level=1)
    go_conditions = [
        (
            "Obtain an auditable receiving-demand profile and replace the constant "
            "demand-bound sensitivities with case-specific finite dispatch, including "
            "supported capture, storage, route, outage, and receiving-capacity inputs."
        ),
        (
            "Resolve or explicitly exclude unsupported source-capture, auxiliary, "
            "network, hydraulic, installed-cost, and lifecycle-emissions terms."
        ),
        (
            "Supply the case storage, route, and finite residual-load evidence "
            "required by the environmental constraint model; retain categorical "
            "exclusion of any unsupported external sink."
        ),
        (
            "Confirm authorship, affiliation, funding, conflicts, originality, "
            "author approval, and a permanent repository identifier."
        ),
        (
            "Repeat the frozen-value, adversarial-review, and clean-build checks "
            "after any new evidence changes the practical endpoint."
        ),
    ]
    for condition in go_conditions:
        document.add_paragraph(condition, style="List Bullet")

    document.add_heading("8. Claims permitted at the current stage", level=1)
    claims = pd.DataFrame(
        [
            {
                "Permitted": (
                    "A theoretical conserved-heat allocation screening bound and "
                    "pathology diagnosis"
                ),
                "Not permitted": (
                    "A practical recommendation to move heat to the selected sink"
                ),
            },
            {
                "Permitted": (
                    "Measured Frontier source potential over valid observations"
                ),
                "Not permitted": (
                    "Measured complete-year Frontier/ORNL delivery from sensitivity "
                    "dispatch"
                ),
            },
            {
                "Permitted": (
                    "Finite repeated-dispatch sensitivities under constant published "
                    "demand bounds and generic storage assumptions"
                ),
                "Not permitted": (
                    "Calibrated or observed annual Frontier/ORNL operation"
                ),
            },
            {
                "Permitted": (
                    "Normalized operational-emissions and generic cost comparators"
                ),
                "Not permitted": (
                    "Frontier project LCOH, NPC, payback, or lifecycle CO2e"
                ),
            },
            {
                "Permitted": "Population-weighted thermal-stress burden",
                "Not permitted": "Mortality, morbidity, or causal health benefit",
            },
        ]
    )
    _add_table(document, claims, "Table 3. Current claim boundary.")

    document.add_heading("9. Recommended next decision", level=1)
    document.add_paragraph(
        "Do not submit the package to Applied Energy at the current evidence state. "
        "The remaining paths are to obtain synchronized demand and an integrated "
        "route, capture, auxiliary, storage, cost, lifecycle, and environmental "
        "design, or to retain the transparent negative result and reconsider the "
        "journal and article type."
    )

    document.add_heading("10. Conclusion", level=1)
    conclusion = document.add_paragraph()
    conclusion.add_run("Final decision: NO-GO.").bold = True
    conclusion.add_run(
        " The mathematical framework, real-case source evidence, technology "
        "comparators, Phase D accounting, environmental exclusions, and Phase F "
        "structural replication, Phase G dispatch sensitivities, constraint "
        "waterfall, mode decision, frozen ledger, and adversarial review constitute "
        "meaningful progress. However, measured finite delivery, case-specific "
        "engineering and economics, case-specific cumulative environmental load, "
        "and submission administration remain unresolved. "
        "The decision should not change to GO until the mandatory gates are closed "
        "with auditable evidence."
    )

    output = SUBMISSION / "nogo_assessment_report.docx"
    document.save(output)
    return output


def _build_editable_figures_tables(context: dict) -> Path:
    presentation = Presentation()
    presentation.slide_width = PptxInches(13.333)
    presentation.slide_height = PptxInches(7.5)
    legends = _figure_legends(context)
    figure_paths = [
        FIGURES / "figure_1_framework.png",
        FIGURES / "figure_2_primary_spatial_transfer.png",
        FIGURES / "figure_3_allocation_mode_comparison.png",
        FIGURES / "figure_4_case_source_and_demand.png",
        FIGURES / "figure_5_constraint_waterfall.png",
        FIGURES / "figure_6_practical_modes.png",
    ]
    for figure_path, legend in zip(figure_paths, legends, strict=True):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        title = slide.shapes.add_textbox(
            PptxInches(0.45),
            PptxInches(0.2),
            PptxInches(12.4),
            PptxInches(0.55),
        )
        title_paragraph = title.text_frame.paragraphs[0]
        title_paragraph.text = legend.split(".", 1)[0]
        title_paragraph.alignment = PP_ALIGN.CENTER
        title_paragraph.runs[0].font.size = PptxPt(24)
        title_paragraph.runs[0].font.bold = True
        with Image.open(figure_path) as image:
            scale = min(11.8 / image.width, 5.4 / image.height)
            width = PptxInches(image.width * scale)
            height = PptxInches(image.height * scale)
        picture = slide.shapes.add_picture(
            str(figure_path),
            (presentation.slide_width - width) // 2,
            PptxInches(0.85),
            width=width,
            height=height,
        )
        picture.top = PptxInches(0.85)
        caption = slide.shapes.add_textbox(
            PptxInches(0.65),
            PptxInches(6.35),
            PptxInches(12.0),
            PptxInches(0.8),
        )
        caption_paragraph = caption.text_frame.paragraphs[0]
        caption_paragraph.text = legend
        caption_paragraph.alignment = PP_ALIGN.CENTER
        caption_paragraph.runs[0].font.size = PptxPt(11)

    for frame, caption_text in _main_tables(context):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        title = slide.shapes.add_textbox(
            PptxInches(0.45),
            PptxInches(0.2),
            PptxInches(12.4),
            PptxInches(0.55),
        )
        title_paragraph = title.text_frame.paragraphs[0]
        title_paragraph.text = caption_text
        title_paragraph.alignment = PP_ALIGN.CENTER
        title_paragraph.runs[0].font.size = PptxPt(20)
        title_paragraph.runs[0].font.bold = True
        table_shape = slide.shapes.add_table(
            len(frame) + 1,
            len(frame.columns),
            PptxInches(0.35),
            PptxInches(1.0),
            PptxInches(12.63),
            PptxInches(5.9),
        )
        table = table_shape.table
        for column_index, column in enumerate(frame.columns):
            table.cell(0, column_index).text = str(column).replace("_", " ").title()
        for row_index, row in enumerate(frame.itertuples(index=False), start=1):
            for column_index, value in enumerate(row):
                text = f"{value:.3g}" if isinstance(value, float) else str(value)
                table.cell(row_index, column_index).text = text
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.text_frame.paragraphs:
                    paragraph.alignment = PP_ALIGN.CENTER
                    for run in paragraph.runs:
                        run.font.size = PptxPt(8)
        for cell in table.rows[0].cells:
            cell.fill.solid()
            cell.fill.fore_color.rgb = PptxRGBColor.from_string("D9EAF2")
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.bold = True

    output = SUBMISSION / "figures_tables_editable.pptx"
    presentation.save(output)
    return output


def _build_submission_checklist() -> Path:
    checklist = """# Applied Energy submission checklist

## Generated and checked

- [x] Full Length Research Article structure
- [x] Abstract, keywords, highlights, tables, and figure legends
- [x] Review manuscript with figures and tables placed inline after first citation
- [x] Figures supplied as separate 300-dpi PNG and vector SVG files
- [x] Editable graphical abstract and figure/table deck supplied as PPTX
- [x] Editable main tables supplied as DOCX
- [x] Standalone NO-GO assessment report with explicit GO-transition gates
- [x] Supplementary methods, tables, and figure legends
- [x] Vancouver references numbered in order of appearance
- [x] Data/code availability and generative-AI declarations
- [x] Conserved-energy, no-relocation, and pathological-result language
- [x] Numerical provenance ledger for generated artifacts and key claims
- [x] Prespecified structural replication across years, products, grids,
      population surfaces, and thermal metrics
- [x] Global-to-real constraint waterfall
- [x] Practical Pareto and mode-decision analysis

## Author confirmation required before submission

- [ ] Confirm author list, order, affiliations, and corresponding-author details
- [ ] Confirm funding statement and grant numbers
- [ ] Confirm competing-interest declaration
- [ ] Confirm originality and approval by all authors
- [x] Activate the public repository and insert its GitHub URL
- [ ] Archive the public repository and insert its permanent DOI
- [ ] Choose subscription or open-access publication route
- [ ] Re-read the live Applied Energy guide immediately before submission
- [ ] Add suggested/opposed reviewers after conflict-of-interest review
- [ ] Confirm every submission-system declaration
- [x] Add an explicit heat-source inventory and documented demand bounds
- [x] Add finite repeated-dispatch sensitivities with explicit energy conservation
- [ ] Replace demand-bound sensitivities with measured receiving demand and case design
- [ ] Calibrate transport/storage losses, efficiency, exergy, cost, and emissions
- [x] Add multi-year/product/grid/population structural uncertainty analyses
- [x] Freeze an immutable revision and pass a clean-checkout full rebuild

## Current scientific gate

**NO-GO FOR SUBMISSION.** The current package combines a theoretical ambient-heat
screening study with partial real-case evidence and scoped structural replication.
Finite repeated-dispatch sensitivities are implemented, but measured receiving
demand and case-specific integrated engineering and economic evidence remain
missing. The constraint waterfall and mode-decision analysis therefore retain
zero supported practical modes. Author declarations are also incomplete.
"""
    output = SUBMISSION / "submission_checklist.md"
    output.write_text(checklist, encoding="utf-8")
    return output


def _copy_figures() -> list[Path]:
    output_directory = SUBMISSION / "figures"
    output_directory.mkdir(parents=True, exist_ok=True)
    for existing in output_directory.iterdir():
        if existing.is_file():
            existing.unlink()
    outputs = []
    for source in sorted(FIGURES.glob("figure_*.png")) + sorted(
        FIGURES.glob("figure_*.svg")
    ):
        output = output_directory / source.name
        shutil.copy2(source, output)
        outputs.append(output)
    for source in sorted(FIGURES.glob("supplementary_figure_*.png")) + sorted(
        FIGURES.glob("supplementary_figure_*.svg")
    ):
        output = output_directory / source.name
        shutil.copy2(source, output)
        outputs.append(output)
    return outputs


def _copy_structural_results() -> Path:
    output = SUBMISSION / "structural_replication.csv"
    shutil.copy2(STRUCTURAL_REPLICATION, output)
    return output


def _build_archive(files: list[Path]) -> Path:
    output = SUBMISSION / "applied_energy_submission_package.zip"
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for file in files:
            item = zipfile.ZipInfo(
                str(file.relative_to(SUBMISSION)),
                date_time=(1980, 1, 1, 0, 0, 0),
            )
            item.compress_type = zipfile.ZIP_DEFLATED
            item.create_system = 3
            item.external_attr = 0o600 << 16
            archive.writestr(item, file.read_bytes())
    return output


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--nogo-report-only",
        action="store_true",
        help="Build only the standalone NO-GO assessment report.",
    )
    args = parser.parse_args()
    SUBMISSION.mkdir(parents=True, exist_ok=True)
    context = _result_context()
    if args.nogo_report_only:
        report = _build_nogo_report(context)
        _normalize_zip_archive(report)
        print(f"Wrote {report.name}.")
        return
    references = _reference_metadata()
    files = [
        _build_manuscript(context, references),
        _build_manuscript(context, references, inline_assets=True),
        _build_supplement(context),
        _build_highlights(context),
        _build_cover_letter(context),
        *_build_graphical_abstract(context),
        _build_editable_figures_tables(context),
        _build_editable_tables(context),
        _build_nogo_report(context),
        _build_submission_checklist(),
        _copy_structural_results(),
        *_copy_figures(),
    ]
    for file in files:
        if file.suffix in {".docx", ".pptx"}:
            _normalize_zip_archive(file)
    archive = _build_archive(files)
    print(f"Wrote {len(files)} submission files and {archive.name}.")


if __name__ == "__main__":
    main()
