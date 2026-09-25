"""Validate the frozen Journal of Cleaner Production submission package."""

from __future__ import annotations

import hashlib
import re
import zipfile
from pathlib import Path

import pandas as pd
from docx import Document
from PIL import Image

if __package__:
    from scripts import build_jcp_submission as jcp
else:
    import build_jcp_submission as jcp

ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "docs" / "JCP_FINAL_QC.md"
CLAIMS = ROOT / "provenance" / "claim_evidence_matrix.csv"
PRACTICAL_MODES = ROOT / "results" / "techno_economic" / "practical_pareto_modes.csv"
PATHWAYS = (
    ROOT / "results" / "environmental_constraints" / "frontier_pathway_assessment.csv"
)


def _document_text(path: Path) -> str:
    document = Document(path)
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    cells = [
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    ]
    return "\n".join([*paragraphs, *cells])


def _check(condition: bool, message: str) -> None:
    if not condition:
        raise ValueError(message)


def main() -> None:
    blocks, references, word_count = jcp._resolve_source()
    resolved_text = "\n".join(text for _, text in blocks)
    source = jcp.MANUSCRIPT_SOURCE.read_text(encoding="utf-8")
    registry = pd.read_csv(jcp.MANUSCRIPT_VALUES).fillna("")
    value_ids = set(registry["value_id"])
    requested_value_ids = set(re.findall(r"\{\{value:([^}]+)\}\}", source))
    _check(requested_value_ids <= value_ids, "Unregistered manuscript value.")
    _check(6000 <= word_count <= 8000, "JCP main-text word count is out of range.")
    _check(len(references) <= 50, "JCP reference limit exceeded.")
    _check(all("&amp;" not in item for item in references), "Encoded reference text.")
    _check(
        all(", ," not in item and ";." not in item for item in references),
        "Malformed reference.",
    )
    _check(
        "{{" not in resolved_text and "[[" not in resolved_text,
        "Unresolved source token.",
    )
    _check(
        not re.search(r"[\u3040-\u30ff\u3400-\u9fff\uff01-\uff60]", resolved_text),
        "Japanese or full-width character detected in manuscript text.",
    )

    figure_numbers = list(map(int, re.findall(r"\[\[FIGURE:(\d+)\]\]", source)))
    table_numbers = list(map(int, re.findall(r"\[\[TABLE:(\d+)\]\]", source)))
    _check(figure_numbers == list(range(1, 7)), "Main figure numbering is incomplete.")
    _check(table_numbers == list(range(1, 6)), "Main table numbering is incomplete.")
    _check(
        "Figures S1-S6" in source,
        "Supplementary figures are not cited as a complete sequence.",
    )
    supplementary_table_citations = []
    for number in map(int, re.findall(r"Table S(\d+)", source)):
        if number <= 4 and number not in supplementary_table_citations:
            supplementary_table_citations.append(number)
    _check(
        supplementary_table_citations[:4] == [1, 2, 3, 4],
        "Supplementary tables S1-S4 are not first cited sequentially.",
    )
    _check(
        "Tables S5-S10" in source,
        "Supplementary tables S5-S10 are not cited.",
    )

    clean_path = jcp.SUBMISSION / "manuscript_jcp.docx"
    inline_path = jcp.SUBMISSION / "manuscript_jcp_inline.docx"
    clean_document = Document(clean_path)
    inline_document = Document(inline_path)
    _check(len(clean_document.inline_shapes) == 0, "Clean manuscript embeds figures.")
    _check(
        len(inline_document.inline_shapes) == 6,
        "Inline manuscript figure count differs.",
    )
    _check(len(clean_document.tables) == 5, "Clean manuscript table count differs.")
    _check(len(inline_document.tables) == 5, "Inline manuscript table count differs.")
    paragraphs = [paragraph.text for paragraph in clean_document.paragraphs]
    abstract_index = paragraphs.index("Abstract")
    abstract_words = len(paragraphs[abstract_index + 1].split())
    _check(abstract_words <= 250, "Abstract exceeds 250 words.")
    keywords = paragraphs[abstract_index + 2].removeprefix("Keywords:").split(";")
    _check(1 <= len(keywords) <= 7, "Keyword count is outside 1-7.")
    _check(
        "CRediT authorship contribution statement" in paragraphs,
        "CRediT authorship contribution statement is missing.",
    )
    with zipfile.ZipFile(clean_path) as document_archive:
        document_xml = document_archive.read("word/document.xml").decode("utf-8")
    _check('w:fill="D9EAF2"' not in document_xml, "Table-cell shading remains.")
    _check(
        '<w:insideV w:val="nil"' in document_xml,
        "Vertical table borders were not disabled.",
    )
    _check("{{" not in _document_text(clean_path), "DOCX contains unresolved values.")
    _check(
        "An 18-hour delay produced" in _document_text(clean_path),
        "Temporal-delay statement is missing its hour unit.",
    )
    boundary_sentence = (
        "A problem does not disappear when it crosses the system boundary."
    )
    _check(
        _document_text(clean_path).count(boundary_sentence) == 1,
        "Boundary sentence count differs in manuscript.",
    )
    cover_text = _document_text(jcp.SUBMISSION / "cover_letter.docx")
    _check(
        cover_text.count(boundary_sentence) == 1,
        "Boundary sentence count differs in cover letter.",
    )
    supplement_text = _document_text(jcp.SUBMISSION / "supplementary_material.docx")
    _check(
        "\nnan\n" not in f"\n{supplement_text}\n",
        "Supplement contains a visible missing-value token.",
    )

    highlights = [
        line.removeprefix("• ").strip()
        for line in (jcp.SUBMISSION / "highlights.txt")
        .read_text(encoding="utf-8")
        .splitlines()
        if line.startswith("• ")
    ]
    _check(3 <= len(highlights) <= 5, "Highlight count is outside 3-5.")
    _check(
        all(len(highlight) <= 85 for highlight in highlights),
        "A highlight exceeds 85 characters.",
    )
    with Image.open(jcp.SUBMISSION / "graphical_abstract.png") as graphic:
        width, height = graphic.size
    _check(
        width >= 1328 and height >= 531,
        "Graphical abstract is below the JCP minimum dimensions.",
    )

    for number in range(1, 7):
        for suffix in (".png", ".svg", ".pdf"):
            figure = jcp.FIGURE_DIR / f"figure_{number}{suffix}"
            _check(figure.stat().st_size > 0, f"Missing or empty {figure.name}.")

    practical_modes = pd.read_csv(PRACTICAL_MODES).fillna("")
    _check(
        practical_modes["eligibility"].eq("CONDITIONAL").sum() == 6,
        "Conditional practical-mode count changed.",
    )
    _check(
        practical_modes["eligibility"].eq("SUPPORTED").sum() == 0,
        "Supported practical-mode count changed.",
    )
    pathways = pd.read_csv(PATHWAYS).fillna("")
    _check(
        not pathways["conclusion"]
        .str.contains("environmentally preferred", case=False)
        .any(),
        "A conditional pathway is described as environmentally preferred.",
    )
    claims = pd.read_csv(CLAIMS).fillna("")
    _check(
        claims["claim_level"].str.startswith("Level 3").any(),
        "Prohibited claim controls are absent.",
    )

    archive_path = jcp.SUBMISSION / "jcp_submission_package.zip"
    expected_members = {
        path.relative_to(jcp.SUBMISSION).as_posix()
        for path in jcp.SUBMISSION.rglob("*")
        if path.is_file() and path != archive_path
    }
    with zipfile.ZipFile(archive_path) as archive:
        archived_members = set(archive.namelist())
        _check(archive.testzip() is None, "ZIP integrity check failed.")
    _check(archived_members == expected_members, "ZIP membership differs from package.")

    checksum = hashlib.sha256(archive_path.read_bytes()).hexdigest()
    author_actions = resolved_text.count("AUTHOR ACTION")
    report = f"""# JCP final quality-control report

## Verdict

**READY FOR AUTHOR ACTION / SUBMISSION.**

The package is technically reproducible and the scoped conservation-aware,
evidence-gated contribution remains within its frozen claim boundaries. This
status does not resolve the explicit author-controlled submission fields.
The prior **WEAK GO** label is superseded by this author-action status.

## Automated checks

- Main-text word count: {word_count} (required target: 6000-8000).
- Verified main references: {len(references)} (maximum: 50).
- Main figures: 6, cited sequentially and supplied separately as PNG, SVG, and
  vector PDF.
- Main tables: 5, cited sequentially and embedded in both manuscript versions.
- Clean manuscript embedded figures: 0.
- Inline manuscript embedded figures: 6.
- Practical modes: 6 conditional, 0 supported.
- Manuscript-value placeholders: all resolved through the value registry.
- Reference order: sequential and matched to first appearance.
- Claim-evidence controls: Level 3 prohibited claims retained.
- Accidental Japanese/full-width-character scan: clear.
- ZIP integrity and exact-membership check: passed.
- ZIP SHA-256: `{checksum}`.

## Author-controlled blockers

The manuscript retains {author_actions} explicit `AUTHOR ACTION` statements.
Before submission, the author must confirm affiliation and postal address,
funding, competing interests, institutional-statement requirements, and the
CRediT contribution statement, and archive the repository with a permanent DOI.

## Interpretation

This verdict is not a claim that a practical thermal pathway is supported. It is a
submission-readiness judgment for a manuscript whose applied endpoint is
explicitly zero evidence-supported modes.
"""
    REPORT.write_text(report, encoding="utf-8")
    print(f"JCP QC passed; wrote {REPORT}")


if __name__ == "__main__":
    main()
